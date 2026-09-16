#!/usr/bin/env python3
"""FAQ v1.2 — reposiciona o precatório no Capítulo 1 (DCF): o foco passa a ser o
direito creditório; precatório fica só como diferenciação pontual. Ajustes espelhados
no Capítulo 3 (transação)."""
import copy
import re
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

SRC = '/home/hermes/agente-rene/work/icms_fix/FAQ_Capacitacao_Comercial_LucroAtivo_v1.1.docx'
OUT = '/home/hermes/agente-rene/work/faq_dcf_fix/FAQ_Capacitacao_Comercial_LucroAtivo_v1.2.docx'

doc = Document(SRC)
body = doc.element.body
blocks = list(body.iterchildren())


def ptext(el):
    return ''.join(t.text or '' for t in el.iter() if t.tag.endswith('}t'))


def preplace(p, old, new, required=True):
    runs = p.runs
    texts = [r.text for r in runs]
    full = ''.join(texts)
    idx = full.find(old)
    if idx < 0:
        if required:
            raise AssertionError(f'NAO ACHOU: {old!r} em {full[:200]!r}')
        return False
    end = idx + len(old)
    pos = 0
    for ri, r in enumerate(runs):
        rlen = len(texts[ri])
        rstart, rend = pos, pos + rlen
        pos = rend
        if rend <= idx or rstart >= end:
            continue
        before = texts[ri][:max(0, idx - rstart)]
        after = texts[ri][max(0, min(rlen, end - rstart)):]
        if rstart <= idx < rend:
            r.text = before + new + after
        else:
            r.text = before + after
    return True


def cell_replace(tbl_el, old, new):
    for pel in tbl_el.iter(qn('w:p')):
        p = Paragraph(pel, doc)
        if old in p.text:
            preplace(p, old, new)
            return True
    raise AssertionError(f'NAO ACHOU NA TABELA: {old!r}')


# ---------- localizar limites ----------
cap1 = cap2 = None
for i, el in enumerate(blocks):
    if el.tag.endswith('}p'):
        t = ptext(el).strip()
        if t.startswith('Capítulo 1') and cap1 is None:
            cap1 = i
        elif t.startswith('Capítulo 2') and cap1 is not None and cap2 is None:
            cap2 = i
assert cap1 == 14 and cap2 == 122, (cap1, cap2)

B = blocks  # atalho; referências de elemento seguem válidas após remoções


def P(i):
    return Paragraph(B[i], doc)


# ---------- templates (do próprio capítulo 1) ----------
TPL = {
    'section': copy.deepcopy(B[18]),   # Heading 1 "BLOCO 1 — CONCEITO"
    'question': copy.deepcopy(B[19]),  # "▸  1. ..."
    'bodyp': copy.deepcopy(B[20]),     # corpo com runs reg+bold
}
q_runs = Paragraph(TPL['question'], doc).runs
b_runs = Paragraph(TPL['bodyp'], doc).runs
# identificar run regular e run bold do corpo (w:b pode vir com val="0")
reg_run = next(r for r in b_runs if not r.bold)
bold_run = next(r for r in b_runs if r.bold)
RT = {
    'q_marker': copy.deepcopy(q_runs[0]._element),
    'q_text': copy.deepcopy(q_runs[1]._element),
    'body_reg': copy.deepcopy(reg_run._element),
    'body_bold': copy.deepcopy(bold_run._element),
}


def set_run_text(rel, text):
    for t in rel.findall(qn('w:t')):
        rel.remove(t)
    t = rel.makeelement(qn('w:t'), {})
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    rel.append(t)
    return rel


def clear_runs(pel):
    for r in pel.findall(qn('w:r')):
        pel.remove(r)


def make_p(template_key, runs_spec):
    el = copy.deepcopy(TPL[template_key])
    clear_runs(el)
    for rkey, text in runs_spec:
        r = copy.deepcopy(RT[rkey])
        set_run_text(r, text)
        el.append(r)
    return el


def question(num, text):
    marker = ptext(RT['q_marker'])  # preserva o marcador exato
    return make_p('question', [('q_marker', marker), ('q_text', f'{num}. {text}')])


def bodyp(spec):
    return make_p('bodyp', [(('body_bold' if b else 'body_reg'), t) for t, b in spec])


def section(title):
    el = copy.deepcopy(TPL['section'])
    runs = el.findall(qn('w:r'))
    keep = runs[0]
    for r in runs[1:]:
        el.remove(r)
    set_run_text(keep, title)
    return el


# ---------- NOVO BLOCO 1 + BLOCO 2 (Q1–Q10) ----------
new_els = []
A = new_els.append

A(section('BLOCO 1 — CONCEITO'))
A(question(1, 'O que é o Direito Creditório Federal (DCF) da Lucro Ativo?'))
A(bodyp([('É a solução da Lucro Ativo de ', 0), ('estruturação e aproveitamento de direitos creditórios federais', 1),
         (' — créditos ', 0), ('líquidos, certos e exigíveis contra a União', 1),
         (', reconhecidos por decisão judicial transitada em julgado — para ', 0),
         ('quitar ou amortizar passivos tributários federais', 1),
         (' da empresa-cliente. A operação combina engenharia jurídica (cessão, habilitação, due diligence), engenharia tributária (transação, uso de prejuízo fiscal) e mitigação de risco (garantias/seguro) para transformar um passivo fiscal em uma operação de liquidação com economia de caixa.', 0)]))
A(question(2, 'O que é um "direito creditório federal"?'))
A(bodyp([('É todo crédito ', 0), ('líquido, certo e exigível contra a União', 1),
         (' (ou suas autarquias/fundações federais), ', 0),
         ('reconhecido por decisão judicial transitada em julgado', 1),
         ('. É o instrumento central do nosso trabalho: a Lucro Ativo estrutura a ', 0), ('cessão', 1),
         (' desses direitos e o seu ', 0), ('aproveitamento administrativo', 1),
         (' para extinguir débitos do cliente. Conforme o desenho, a operação também mobiliza ', 0),
         ('créditos próprios do contribuinte', 1),
         (' — prejuízo fiscal (PF) e base de cálculo negativa (BCN) de CSLL — pelo caminho da transação.', 0)]))
A(question(3, 'Direito creditório é o mesmo que precatório?'))
A(bodyp([('Não — e é só para marcar essa diferença que o precatório aparece neste FAQ.', 1),
         (' Precatório é a requisição de pagamento expedida pela Justiça contra a Fazenda Pública (art. 100 da CF), que entra na fila de pagamentos da União. O nosso trabalho é com o ', 0),
         ('direito creditório em si', 1),
         (' — o crédito reconhecido judicialmente, adquirido por cessão e aproveitado como ', 0),
         ('instrumento de quitação', 1),
         (' de débitos, sem depender de fila de pagamento. Na reunião, cite o precatório apenas de forma pontual e superficial, para explicar o que a nossa operação ', 0),
         ('não é', 1),
         ('. Se o cliente quiser aprofundar o regime de precatórios, leve à área técnica.', 0)]))
A(question(4, 'De quem é o crédito? Qual o papel da Lucro Ativo?'))
A(bodyp([('A Lucro Ativo é ', 0), ('cessionária formal', 1),
         (' dos direitos creditórios que estrutura (cessão nos termos do Código Civil, arts. 286–298) e ', 0),
         ('conduz a operação integralmente', 1),
         (': due diligence do crédito, formalização da cessão, habilitação perante o órgão competente e acompanhamento até a baixa do débito. O cliente não "compra um crédito no mercado" por conta própria — ele adere a uma operação estruturada e conduzida pela Lucro Ativo.', 0)]))
A(question(5, 'O que significa "aproveitamento administrativo" do direito creditório?'))
A(bodyp([('Significa que o crédito é utilizado ', 0), ('perante o órgão competente', 1),
         (' — pela ', 0), ('habilitação', 1), (' e pelo ', 0), ('encontro de contas', 1),
         (' — para ', 0), ('extinguir ou reduzir o débito', 1),
         (' do cliente. Não se trata de propor nova ação judicial, nem de aguardar o recebimento do crédito em espécie: o direito creditório funciona como ', 0),
         ('meio de quitação', 1), (' do passivo.', 0)]))
A(question(6, 'O que é "crédito homologado" (PF/BCN) e qual a diferença para o direito creditório?'))
A(bodyp([('"Crédito homologado" é o crédito tributário do ', 0), ('próprio contribuinte', 1),
         (' reconhecido/aceito pela Receita Federal — tipicamente ', 0),
         ('prejuízo fiscal e base de cálculo negativa de CSLL', 1),
         (', apurados e declarados ao Fisco, ou créditos de indébito. A diferença central: o ', 0),
         ('direito creditório', 1),
         (' é crédito judicial, pode ser de terceiro (adquirido por cessão) e serve para pagar débitos inscritos em dívida ativa por encontro de contas; já ', 0),
         ('PF/BCN', 1),
         (' é crédito do próprio contribuinte (ou de coligada/controlada), usado dentro da transação tributária para abater o saldo. São caminhos jurídicos distintos, com bases legais distintas — nunca os apresente como a mesma coisa.', 0)]))

A(section('BLOCO 2 — BASE LEGAL'))
A(question(7, 'Qual a base legal da cessão do direito creditório?'))
A(bodyp([('O ', 0), ('Código Civil, arts. 286 a 298', 1),
         (': o art. 286 permite ao credor ceder seu crédito, salvo vedação legal, da natureza da obrigação ou de convenção; o art. 287 estende a cessão aos acessórios do crédito; o art. 288 exige instrumento adequado para eficácia contra terceiros; e o art. 290 condiciona a eficácia perante o devedor à notificação (ou ciência declarada). É esse arcabouço que dá segurança jurídica à aquisição do direito creditório por cessão — e o motivo de a ', 0),
         ('cadeia dominial', 1), (' documentada ser inegociável.', 0)]))
A(question(8, 'Qual a base do uso de créditos judiciais para quitar débitos federais?'))
A(bodyp([('As âncoras que o executivo cita (sem detalhar o rito): o ', 0),
         ('art. 100, § 11, da CF', 1),
         (' (redação da EC 113/2021) autoriza o credor a ofertar créditos reconhecidos em decisão transitada em julgado para quitar débitos parcelados ou inscritos em dívida ativa do ente devedor — e, para a União, o mecanismo é ', 0),
         ('autoaplicável', 1),
         ('. A regulamentação vem do ', 0), ('Decreto nº 11.249/2022', 1),
         (' (procedimento de oferta e encontro de contas) e da ', 0), ('Portaria PGFN nº 10.826/2022', 1),
         (' (requisitos, documentação e uso por cessionário/terceiro, com a cadeia dominial do crédito até o último cessionário). ', 0),
         ('O enquadramento exato de cada operação é definido pela área técnica — (PARÂMETRO A VALIDAR)', 1),
         ('; o papel do executivo é citar as âncoras com segurança, não opinar sobre o rito.', 0)]))
A(question(9, 'Qual a base legal do uso de crédito próprio (PF/BCN) para quitar tributo?'))
A(bodyp([('A ', 0), ('transação tributária', 1), (': ', 0), ('Lei nº 13.988/2020, art. 11', 1),
         (', alterada pela ', 0), ('Lei nº 14.375/2022', 1),
         (', que autoriza o uso de créditos de ', 0), ('prejuízo fiscal (PF)', 1), (' e ', 0),
         ('base de cálculo negativa de CSLL (BCN)', 1),
         (' para abater até 70% do saldo remanescente após os descontos da transação. O art. 11, § 7º admite PF/BCN do próprio contribuinte, de controladora, controlada, ou de empresas controladas pela mesma pessoa jurídica, apurados e declarados à Receita.', 0)]))
A(question(10, 'E a compensação comum do art. 74 da Lei 9.430/1996 — cabe aqui?'))
A(bodyp([('Atenção — esse é o ponto que mais gera confusão.', 1),
         (' O ', 0), ('art. 74 da Lei 9.430/1996', 1),
         (' permite ao contribuinte compensar tributos federais com ', 0), ('créditos próprios', 1),
         (' (inclusive de decisão judicial transitada em julgado), via PER/DCOMP. Porém, a lei ', 0),
         ('veda a compensação com créditos de terceiros', 1),
         (' (art. 74, § 12), e a Receita Federal alerta publicamente contra ofertas de "créditos de terceiros" para compensação. Por isso o direito creditório adquirido de terceiro ', 0),
         ('NÃO entra por compensação do art. 74', 1),
         (' — ele entra pelo caminho do art. 100, § 11, da CF (encontro de contas com dívida ativa), que é juridicamente distinto. Saber essa diferença é o que separa um produto sério de uma fraude.', 0)]))

# ---------- trocas cirúrgicas (capítulo 1, blocos 40+) ----------
preplace(P(44), 'precatório adquirido (para dívida ativa', 'direito creditório adquirido (para dívida ativa')
preplace(P(45), 'quando via precatório,', 'quando via direito creditório de terceiro,')
preplace(P(47), 'dívida ativa (precatório)', 'dívida ativa (direito creditório)')
preplace(P(50), '(precatório/direito creditório)', '(direito creditório)')
preplace(P(55), 'caminho do precatório/encontro de contas', 'caminho do encontro de contas')
preplace(P(57), 'uso de precatório para dívida ativa', 'uso de direito creditório para dívida ativa')
preplace(P(57), 'pelo caminho do precatório', 'pelo caminho do direito creditório')
preplace(P(58), 'PRECATÓRIO DE TERCEIRO', 'DIREITO CREDITÓRIO DE TERCEIRO')
preplace(P(59), 'usar precatório de terceiro', 'usar direito creditório de terceiro')
cell_replace(B[60], 'Precatório de terceiro (adquirido)', 'Direito creditório de terceiro (adquirido)')
preplace(P(68), 'higidez do precatório', 'higidez do crédito')
preplace(P(70), 'Precatório ainda em discussão', 'Crédito ainda em discussão')
preplace(P(73), 'aquisição do precatório', 'aquisição do direito creditório')
preplace(P(76), 'autoriza usar precatório para quitar débito', 'autoriza a oferta de créditos reconhecidos judicialmente para quitar débitos')
preplace(P(77), 'E se o precatório não for pago pela União?', 'E se a União não pagar esse crédito?')
preplace(P(78), 'o precatório não é comprado para esperar o pagamento da fila',
         'o direito creditório não é adquirido para se aguardar pagamento (não é aposta em fila de precatórios)')
preplace(P(84), 'deságio do precatório', 'deságio do direito creditório')
cell_replace(B[87], 'Adquire precatório federal de', 'Adquire direito creditório federal de')
cell_replace(B[87], 'Oferece o precatório em encontro de contas', 'Oferece o crédito em encontro de contas')
preplace(P(92), 'garante pagamento de precatório na fila nem', 'é aposta em fila de precatórios nem promessa de')

# glossário: nova entrada "Direito creditório federal" antes de "Precatório:"
gloss_el = copy.deepcopy(B[98])
gp = Paragraph(gloss_el, doc)
gruns = gp.runs
gruns[0].text = 'Direito creditório federal: '
rest = ('crédito líquido, certo e exigível contra a União, reconhecido por decisão judicial '
        'transitada em julgado — o instrumento central da operação, adquirido por cessão e '
        'aproveitado administrativamente para quitação de débitos.')
gruns[1].text = rest
for r in gruns[2:]:
    r.text = ''
B[98].addprevious(gloss_el)

preplace(P(98), 'acima do teto de RPV.',
         'acima do teto de pequeno valor. Citado neste FAQ apenas para diferenciação — não é o instrumento da operação.')
body.remove(B[99])  # entrada RPV (detalhe do regime de precatórios — dispensável)
preplace(P(104), 'o precatório ofertado é abatido', 'o crédito ofertado é abatido')
preplace(P(112), 'regime de precatórios', 'pagamentos devidos pela Fazenda Pública em virtude de decisão judicial')
preplace(P(114), 'teto de precatórios até 2026', 'regime de pagamento dos débitos judiciais da União até 2026')
preplace(P(116), 'uso de precatórios para dívida ativa', 'oferta de créditos judiciais para amortização de dívida ativa')

# ---------- capítulo 3 ----------
preplace(P(246), 'e precatórios federais —', 'e demais direitos creditórios federais admitidos —')
preplace(P(250), 'quitar com PF/BCN/precatório', 'quitar com PF/BCN e outros créditos admitidos')
preplace(P(287), 'e/ou precatórios federais', 'e/ou direitos creditórios federais admitidos')
preplace(P(288), 'USO DE CRÉDITOS (PF, BCN, PRECATÓRIO)', 'USO DE CRÉDITOS (PF, BCN E DIREITOS CREDITÓRIOS)')
preplace(P(295), 'E precatórios federais?', 'E direitos creditórios federais — inclusive precatórios?')
preplace(P(296), 'Precatórios federais (próprios ou adquiridos) podem',
         'Direitos creditórios federais (próprios ou adquiridos — categoria que inclui os precatórios) podem')
preplace(P(305), 'e/ou precatórios', 'e/ou direitos creditórios')
preplace(P(351), 'Precatório federal:', 'Direito creditório federal:')
preplace(P(351), 'título representativo de crédito reconhecido judicialmente contra a União,',
         'crédito reconhecido judicialmente contra a União (categoria que inclui os precatórios),')
preplace(P(365), 'crédito de terceiro (precatório)', 'crédito de terceiro (direito creditório adquirido)')

# ---------- aplicar rebuild dos blocos 18–39 ----------
anchor = B[40]  # BLOCO 3 — heading (fica)
for el in B[18:40]:
    body.remove(el)
for el in new_els:
    anchor.addprevious(el)

# ---------- versão v1.1 -> v1.2 (corpo + headers/footers) ----------
def bump_version(paragraphs):
    n = 0
    for p in paragraphs:
        if 'v1.1' in p.text or 'V1.1' in p.text or '1.1' in p.text:
            for r in p.runs:
                if '1.1' in r.text:
                    r.text = r.text.replace('v1.1', 'v1.2').replace('V1.1', 'V1.2').replace('1.1', '1.2')
                    n += 1
    return n

nv = bump_version(doc.paragraphs)
for sec in doc.sections:
    for part in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer):
        if part is not None:
            nv += bump_version(part.paragraphs)
print('versoes atualizadas:', nv)

doc.save(OUT)
print('salvo', OUT)

# ---------- verificação final ----------
check = Document(OUT)
hits = []
for p in check.paragraphs:
    if re.search(r'precat', p.text, re.I):
        hits.append(p.text.strip()[:160])
for tbl in check.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if re.search(r'precat', p.text, re.I):
                    hits.append('[TBL] ' + p.text.strip()[:160])
print(f'\n==== menções restantes a precatório: {len(hits)} ====')
for h in hits:
    print('-', h)
