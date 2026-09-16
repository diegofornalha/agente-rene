#!/usr/bin/env python3
"""Reescreve o Capítulo 2 do FAQ (Recuperação -> Compensação de ICMS com créditos acumulados)."""
import copy
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

SRC = 'FAQ.docx'
OUT = 'FAQ_Capacitacao_Comercial_LucroAtivo_v1.1.docx'

doc = Document(SRC)
body = doc.element.body
blocks = [c for c in body.iterchildren()]


def ptext(el):
    return ' '.join(t.text or '' for t in el.iter() if t.tag.endswith('}t')).strip()

# localizar limites do capítulo 2
start = end = None
for i, el in enumerate(blocks):
    tag = el.tag.split('}')[1]
    if tag == 'p':
        t = ptext(el)
        if t.startswith('Capítulo 2') and start is None:
            start = i
        elif t.startswith('Capítulo 3') and start is not None:
            end = i
            break
assert start is not None and end is not None, (start, end)

# templates (deepcopy antes de deletar)
TPL = {
    'kicker': copy.deepcopy(blocks[start + 1]),
    'howto': copy.deepcopy(blocks[start + 2]),   # tbl COMO USAR
    'gap': copy.deepcopy(blocks[start + 3]),      # p vazio
    'section': copy.deepcopy(blocks[start + 4]),  # Heading 1 SEÇÃO
    'question': copy.deepcopy(blocks[start + 5]),
    'bodyp': copy.deepcopy(blocks[start + 6]),
}
# bullet, atencao, types, exemplo — achar por conteúdo/estilo
for i in range(start, end):
    el = blocks[i]
    tag = el.tag.split('}')[1]
    t = ptext(el)
    if tag == 'p' and 'bullet' not in TPL:
        p = Paragraph(el, doc)
        if p.style.name == 'List Bullet':
            TPL['bullet'] = copy.deepcopy(el)
    if tag == 'tbl':
        if t.startswith('⬧ Atenção') and 'atencao' not in TPL:
            TPL['atencao'] = copy.deepcopy(el)
        if t.startswith('Modalidade'):
            TPL['types'] = copy.deepcopy(el)
        if t.startswith('⬧ EXEMPLO'):
            TPL['exemplo'] = copy.deepcopy(el)
assert all(k in TPL for k in ('bullet', 'atencao', 'types', 'exemplo')), TPL.keys()

# run templates: (marker gold, texto navy) da pergunta; (reg, bold) do corpo; do bullet
q_runs = Paragraph(TPL['question'], doc).runs
RT = {
    'q_marker': copy.deepcopy(q_runs[0]._element),
    'q_text': copy.deepcopy(q_runs[1]._element),
}
b_runs = Paragraph(TPL['bodyp'], doc).runs
RT['body_reg'] = copy.deepcopy(b_runs[0]._element)   # regular Inter dark
RT['body_bold'] = copy.deepcopy(b_runs[1]._element)  # bold Inter navy
bl_runs = Paragraph(TPL['bullet'], doc).runs
RT['bul_bold'] = copy.deepcopy(bl_runs[0]._element)
RT['bul_reg'] = copy.deepcopy(bl_runs[1]._element)
k_runs = Paragraph(TPL['kicker'], doc).runs
RT['kicker'] = copy.deepcopy(k_runs[0]._element)


def set_run_text(rel, text):
    for t in rel.findall(qn('w:t')):
        rel.remove(t)
    t = rel.makeelement(qn('w:t'), {})
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    rel.append(t)
    return rel


def clear_runs(pel):
    for r in pel.findall(qn('w:r')):
        pel.remove(r)


def make_p(template_key, runs_spec):
    """runs_spec: list of (run_template_key, text)"""
    el = copy.deepcopy(TPL[template_key])
    clear_runs(el)
    for rkey, text in runs_spec:
        r = copy.deepcopy(RT[rkey])
        set_run_text(r, text)
        el.append(r)
    return el


def question(num, text):
    return make_p('question', [('q_marker', '▸  '), ('q_text', f'{num}. {text}')])


def bodyp(spec):
    """spec: list of (text, bold_bool)"""
    return make_p('bodyp', [(('body_bold' if b else 'body_reg'), t) for t, b in spec])


def bullet(spec):
    return make_p('bullet', [(('bul_bold' if b else 'bul_reg'), t) for t, b in spec])


def section(title):
    el = copy.deepcopy(TPL['section'])
    # heading: manter estilo, trocar texto de todos os runs -> um só
    runs = el.findall(qn('w:r'))
    keep = runs[0]
    for r in runs[1:]:
        el.remove(r)
    set_run_text(keep, title)
    return el


def kicker(text):
    return make_p('kicker', [('kicker', text)])


def gap():
    return copy.deepcopy(TPL['gap'])


def box(template_key, lead, rest):
    """clona tabela-box e reescreve o parágrafo de corpo (2º parágrafo da célula)."""
    el = copy.deepcopy(TPL[template_key])
    cell_ps = [p for p in el.iter(qn('w:p'))]
    # último parágrafo com texto = corpo
    body_ps = [p for p in cell_ps if ptext(p)]
    target = body_ps[-1]
    runs = target.findall(qn('w:r'))
    assert len(runs) >= 1
    r0 = runs[0]
    set_run_text(r0, lead)
    if len(runs) >= 2:
        r1 = runs[1]
        set_run_text(r1, rest)
        for r in runs[2:]:
            target.remove(r)
    else:
        r1 = copy.deepcopy(r0)
        # tirar negrito do clone se possível
        rpr = r1.find(qn('w:rPr'))
        if rpr is not None:
            for btag in rpr.findall(qn('w:b')):
                rpr.remove(btag)
        set_run_text(r1, rest)
        target.append(r1)
    return el


def types_table(rows):
    """clona a tabela 4x4 e reescreve os textos das células."""
    el = copy.deepcopy(TPL['types'])
    trs = el.findall(qn('w:tr'))
    assert len(trs) == len(rows), (len(trs), len(rows))
    for tr, rowvals in zip(trs, rows):
        tcs = tr.findall(qn('w:tc'))
        assert len(tcs) == len(rowvals), (len(tcs), rowvals)
        for tc, val in zip(tcs, rowvals):
            ps = tc.findall(qn('w:p'))
            first = ps[0]
            runs = first.findall(qn('w:r'))
            if runs:
                set_run_text(runs[0], val)
                for r in runs[1:]:
                    first.remove(r)
            for p in ps[1:]:
                tc.remove(p)
    return el


# ---------- CONTEÚDO NOVO ----------
new_els = []
A = new_els.append

# título do capítulo: reaproveitar o el original (blocks[start]) — retexto depois.
A(kicker('PRODUTO: COMPENSAÇÃO DO ICMS MENSAL COM CRÉDITOS ACUMULADOS TRANSFERIDOS DE TERCEIROS (O CLIENTE É A CESSIONÁRIA)'))
A(copy.deepcopy(TPL['howto']))
A(gap())

A(section('SEÇÃO 1 — CONCEITO'))
A(question(1, 'O que é, em uma frase, o produto de ICMS-SP?'))
A(bodyp([('É a ', 0), ('compensação do ICMS mensal a recolher', 1),
         (' da empresa com ', 0), ('créditos acumulados de ICMS de terceiros', 1),
         (', transferidos para a conta da empresa junto à SEFAZ/SP: o crédito é lançado na apuração (GIA e SPED), ', 0),
         ('reduz o imposto do mês na fonte', 1),
         (' e o cliente paga pelo crédito utilizado com deságio — economia direta, recorrente, mês a mês.', 0)]))
A(question(2, 'Por que NÃO chamamos de "recuperação de ICMS"?'))
A(bodyp([('Recuperação', 1),
         (' é reaver imposto próprio pago a maior no passado (retroativo) — outra tese, com outro rito. O nosso produto atua no ', 0),
         ('imposto corrente', 1),
         (': a empresa deixa de desembolsar o valor integral do ICMS do mês porque compensa a maior parte com crédito transferido. Os termos corretos são ', 0),
         ('compensação e aproveitamento de créditos acumulados', 1),
         ('. Usar "recuperação" cria expectativa errada (de restituição retroativa) e descreve o que o produto não é.', 0)]))
A(question(3, 'De onde vêm os créditos utilizados?'))
A(bodyp([('De ', 0), ('saldos acumulados por empresas incorporadas por multinacionais', 1),
         (' que, ao levantarem seus ativos, absorveram os respectivos créditos — originados das operações de compra, venda, prestação de serviços e importação. Esses créditos encontram-se ', 0),
         ('alocados em conta eletrônica junto à SEFAZ/SP, devidamente habilitados e escriturados digitalmente', 1),
         (', passíveis de transferência para compensação do saldo de ICMS das empresas destinatárias.', 0)]))
A(question(4, 'O que é "crédito acumulado" e por que ele existe?'))
A(bodyp([('O ICMS é ', 0), ('não-cumulativo', 1),
         (': credita-se o imposto das entradas e debita-se o das saídas. Quando, de forma estrutural, os créditos superam os débitos, forma-se um ', 0),
         ('saldo credor', 1),
         (' que a operação não absorve. Nas hipóteses legais (RICMS-SP, art. 71), esse saldo é reconhecido como ', 0),
         ('crédito acumulado', 1),
         (' — apropriável e transferível. As causas típicas no detentor do crédito:', 0)]))
A(bullet([('Exportação', 1), (' — a saída para o exterior é imune/desonerada, mas a lei garante a manutenção do crédito das entradas. O crédito entra, não há débito na saída, acumula.', 0)]))
A(bullet([('Diferimento / suspensão', 1), (' — o imposto é postergado para etapa posterior da cadeia; o estabelecimento faz saídas sem débito (ou com débito reduzido), mas se credita das entradas.', 0)]))
A(bullet([('Saída com alíquota menor que a de entrada', 1), (' — compra a 18% e vende a 7% ou 12% (interestadual, produtos com carga reduzida etc.). A diferença vira saldo credor estrutural.', 0)]))
A(bullet([('Isenção/não-incidência com manutenção de crédito', 1), (' — saídas desoneradas em que a lei permite manter o crédito da entrada.', 0)]))

A(section('SEÇÃO 2 — BASE LEGAL (verificada)'))
A(question(5, 'Qual a base legal do produto? (o executivo precisa citar com segurança)'))
A(bullet([('Constituição Federal, art. 155, §2º:', 1), (' institui a não-cumulatividade do ICMS — a raiz constitucional do direito ao crédito.', 0)]))
A(bullet([('Lei Complementar nº 87/1996 (Lei Kandir):', 1), (' norma nacional do ICMS. Arts. 19 a 23 (direito ao crédito) e art. 25 (saldos credores e transferência de crédito acumulado a terceiros).', 0)]))
A(bullet([('RICMS-SP — Decreto nº 45.490/2000, arts. 71 a 84:', 1), (' o regime do crédito acumulado — hipóteses geradoras (art. 71), geração/apropriação e formas de utilização, incluindo a transferência a outros contribuintes.', 0)]))
A(bullet([('RICMS-SP, art. 99 — o lançamento:', 1), (' fundamento citado no material-base da operação para o lançamento do crédito transferido no campo "Outros Créditos" da GIA e no SPED, com a compensação homologada junto à SEFAZ/SP.', 0)]))
A(bullet([('Portaria SRE nº 65/2023:', 1), (' disciplina a apropriação e a utilização do crédito acumulado e o sistema e-CredAc — a "conta eletrônica" onde os créditos ficam alocados e são transferidos (substituiu a Portaria CAT nº 26/2010). Desde 1º/jan/2026, a fonte de dados dos pedidos passou da GIA para a EFD (SPED Fiscal).', 0)]))
A(copy.deepcopy(TPL['atencao']))  # observação de precisão — texto original serve
A(gap())

A(section('SEÇÃO 3 — NÃO CONFUNDIR (papéis e teses vizinhas)'))
A(question(6, 'Qual a diferença entre o nosso produto, a apropriação de crédito próprio e a "recuperação"?'))
A(types_table([
    ['Modalidade', 'O que é', 'De quem é o crédito', 'Papel no nosso produto'],
    ['Compensação com crédito transferido', 'Abater o ICMS mensal com crédito acumulado recebido de terceiro', 'Do detentor (cedente), habilitado na SEFAZ/SP', 'É o produto — o nosso cliente é a cessionária'],
    ['Apropriação de crédito próprio (e-CredAc)', 'Empresa que acumula saldo credor monetiza o próprio crédito', 'Da própria empresa', 'É o papel do detentor/cedente — não é o nosso cliente'],
    ['Recuperação (ressarcimento de ST, crédito extemporâneo)', 'Reaver imposto próprio pago a maior, retroativamente', 'Da própria empresa', 'Outra tese — fora do escopo deste produto'],
]))
A(gap())
A(bodyp([('Regra de bolso para o executivo:', 1),
         (' o nosso cliente é quem paga muito ICMS todo mês — não quem acumula crédito. Quem acumula é o detentor, que cede o crédito com deságio. Se o prospect tem saldo credor próprio, ele não é comprador deste produto (pode, no limite, ser um futuro detentor — leve à área técnica).', 0)]))

A(section('SEÇÃO 4 — COMO A OPERAÇÃO FUNCIONA (passo a passo)'))
A(question(7, 'Como funciona a operação, do início ao fim?'))
A(bullet([('Qualificação', 1), (' — apuração do ICMS mensal do cliente (GIA/SPED) e dimensionamento do crédito necessário.', 0)]))
A(bullet([('Crédito habilitado', 1), (' — os créditos do detentor já estão alocados e habilitados em conta eletrônica junto à SEFAZ/SP, prontos para transferência.', 0)]))
A(bullet([('Lançamento', 1), (' — solicitamos a GIA no formato PRF, executamos o lançamento no campo "Outros Créditos" e devolvemos o arquivo já transmitido; a escrituração é espelhada no SPED.', 0)]))
A(bullet([('Compensação homologada', 1), (' — a compensação é homologada junto à SEFAZ/SP; o abatimento ocorre na fonte e fica registrado na GIA, no SPED e na Conta Fiscal.', 0)]))
A(bullet([('Constatação', 1), (' — a cessionária confere o lançamento pela conta eletrônica e a Conta Fiscal do mês (guias quitadas, saldo do período zerado).', 0)]))
A(bullet([('Saldo remanescente', 1), (' — a parcela não compensada é recolhida normalmente por DARE.', 0)]))
A(bullet([('Pagamento do crédito', 1), (' — o cliente paga pelo crédito utilizado (com deságio) somente após a constatação da conta fiscal zerada.', 0)]))
A(question(8, 'Qual o papel da Lucro Ativo em cada etapa?'))
A(bodyp([('Condução ', 0), ('integral', 1),
         (' do procedimento: dimensionamento, lançamento e transmissão da GIA, escrituração no SPED, habilitação e compensação na conta fiscal e comprovação documental (recibos de entrega da EFD, conta eletrônica, Conta Fiscal). E o compromisso de compliance do material-base: ', 0),
         ('responsabilidade do escritório pelos 5 anos seguintes a cada compensação', 1),
         (' — qualquer questão a retificar junto à SEFAZ permanece sob a responsabilidade do nosso compliance.', 0)]))

A(section('SEÇÃO 5 — SISTEMAS E TERMOS (glossário operacional)'))
A(question(9, 'O que é a GIA e o que é o formato PRF?'))
A(bodyp([('GIA (Guia de Informação e Apuração do ICMS)', 1),
         (' é a declaração mensal do contribuinte paulista à SEFAZ-SP, que resume a apuração (débitos, créditos e saldos). ', 0),
         ('PRF', 1),
         (' é o formato do arquivo da GIA que o cliente nos envia: executamos o lançamento em "Outros Créditos" e devolvemos o arquivo já transmitido.', 0)]))
A(question(10, 'O que é o SPED Fiscal?'))
A(bodyp([('SPED Fiscal (EFD-ICMS/IPI)', 1),
         (' é a escrituração fiscal digital da empresa. É nele que a compensação é espelhada e habilitada; o recibo de entrega da EFD integra a comprovação documental da operação.', 0)]))
A(question(11, 'O que é o e-CredAc / a "conta eletrônica"?'))
A(bodyp([('e-CredAc (Sistema Eletrônico de Gerenciamento do Crédito Acumulado)', 1),
         (' é o sistema da SEFAZ-SP onde o crédito acumulado é controlado, alocado e transferido eletronicamente (Portaria SRE 65/2023). É por ele que a cessionária constata que o crédito foi efetivamente lançado e aceito na compensação.', 0)]))
A(question(12, 'O que são a Conta Fiscal e a DARE?'))
A(bodyp([('Conta Fiscal', 1),
         (' é o extrato/conta corrente do contribuinte na SEFAZ-SP — é nela que se constata, mês a mês, as guias quitadas e o saldo zerado antes do pagamento do crédito. ', 0),
         ('DARE', 1),
         (' é a guia de recolhimento do saldo remanescente não compensado.', 0)]))

A(section('SEÇÃO 6 — NÚMEROS DA OPERAÇÃO (premissas do material-base)'))
A(question(13, 'Quais são os números de referência da operação?'))
A(bullet([('Utilização do crédito:', 1), (' entre 80% e 90% do ICMS apurado no mês; o remanescente é recolhido pela empresa por DARE.', 0)]))
A(bullet([('Economia:', 1), (' aproximadamente 30% ao mês, obtida com o deságio na aquisição do crédito.', 0)]))
A(bullet([('Honorários:', 1), (' não cobramos honorários do cliente para a execução do trabalho — as despesas do procedimento são arcadas pelo detentor dos créditos.', 0)]))
A(bullet([('Momento do pagamento:', 1), (' o pagamento referente ao crédito utilizado ocorre somente após a constatação da conta fiscal zerada.', 0)]))
A(box('atencao',
      'Estes números vêm do material comercial vigente (Apresentação ICMS-SP, base da operação). ',
      'Em proposta individual, reconfirme com a área técnica antes de personalizar — condição por cliente é (PARÂMETRO A VALIDAR).'))
A(gap())

A(section('SEÇÃO 7 — QUEM SE QUALIFICA (perfil ideal)'))
A(question(14, 'Qual o perfil ideal de cliente para este produto?'))
A(bullet([('Contribuintes de ICMS em São Paulo', 1), (' com imposto a recolher relevante e recorrente (indústria, distribuição, comércio, importação).', 0)]))
A(bullet([('Lucro Real ou Presumido', 1), (', com apuração organizada (GIA/SPED em dia) e situação cadastral regular.', 0)]))
A(bullet([('Caixa pressionado pelo ICMS mensal', 1), (' — o produto devolve fôlego imediato, sem litígio e sem mudar a operação do cliente.', 0)]))
A(question(15, 'Que "gatilhos" o executivo deve escutar na conversa para identificar oportunidade?'))
A(bodyp([('"Pago muito ICMS todo mês"; "o ICMS consome meu caixa"; "minha margem some no imposto"; "não tenho crédito para abater o débito". ', 0),
         ('Qualquer um desses é sinal de qualificação', 1),
         (' — o passo seguinte é pedir a apuração (GIA/SPED) do último mês para dimensionar.', 0)]))

A(section('SEÇÃO 8 — OBJEÇÕES COMUNS E RESPOSTAS'))
A(question(16, '"Isso é legal? É seguro?"'))
A(bodyp([('Totalmente. A ', 0), ('transferência de crédito acumulado é prevista pelo próprio regulamento paulista', 1),
         (' (RICMS-SP, arts. 71-84; Lei Kandir, art. 25). Tudo ocorre nos sistemas oficiais, com homologação da SEFAZ/SP, e fica registrado na GIA, no SPED e na Conta Fiscal. Não é liminar nem tese em disputa: é mecânica regulamentada — e o contrato prevê a responsabilidade do escritório por 5 anos após cada compensação.', 0)]))
A(question(17, '"Por que meu contador nunca me ofereceu isso?"'))
A(bodyp([('O contador cuida da rotina fiscal. O ', 0), ('acesso a créditos habilitados de detentores', 1),
         (' e a condução da transferência/compensação é especialidade de mercado — depende de relacionamento com os detentores e de engenharia documental própria. É um trabalho complementar ao da contabilidade (que, aliás, participa enviando a GIA/SPED), não concorrente.', 0)]))
A(question(18, '"E se a SEFAZ questionar depois?"'))
A(bodyp([('A compensação passa por ', 0), ('homologação', 1),
         (' e a comprovação fica registrada (recibos da EFD, conta eletrônica, Conta Fiscal). Havendo qualquer questão a retificar, o contrato prevê a ', 0),
         ('responsabilidade do escritório pelos 5 anos seguintes', 1),
         (' a cada compensação — permanece sob o nosso compliance, não sobra para o cliente resolver sozinho.', 0)]))
A(question(19, '"Quanto custa?"'))
A(bodyp([('Sem honorários do cliente', 1),
         (': paga-se apenas o crédito efetivamente utilizado, com deságio, e somente após a conta fiscal zerada. As despesas do procedimento são arcadas pelo detentor dos créditos. Condição específica por cliente: (PARÂMETRO A VALIDAR).', 0)]))
A(question(20, '"Parece bom demais para ser verdade."'))
A(bodyp([('A economia tem lógica de mercado: o detentor tem crédito parado (baixa liquidez) e prefere monetizá-lo com desconto; o cliente tem imposto certo todo mês e adquire esse crédito por menos do que pagaria ao Fisco. O ', 0),
         ('deságio é a ponte', 1),
         (' entre os dois — uma mecânica prevista em regulamento, com a SEFAZ/SP homologando cada etapa e todo o rastro documental disponível.', 0)]))

A(section('SEÇÃO 9 — EXEMPLO PRÁTICO (ilustrativo)'))
A(question(21, 'Um exemplo numérico para o executivo entender a mecânica.'))
A(box('exemplo',
      'Exemplo meramente ilustrativo — números fictícios, não é promessa de resultado. ',
      'Indústria paulista apura ICMS de R$ 500.000 no mês. Com a operação: lançamento de crédito transferido cobrindo 85% do imposto (R$ 425.000) e recolhimento do remanescente por DARE (R$ 75.000). Pelo crédito utilizado, a empresa paga ao detentor o valor com deságio — desembolsando menos do que os R$ 425.000 que pagaria ao Fisco. Resultado: o ICMS do mês sai por fração do valor original e a economia se repete enquanto a operação durar (percentuais exatos por cliente: PARÂMETRO A VALIDAR).'))
A(gap())

A(section('SEÇÃO 10 — LIMITES E RISCOS'))
A(question(22, 'O que pode travar ou atrasar a operação?'))
A(bullet([('Apuração inconsistente', 1), (' — divergência entre GIA, SPED e Conta Fiscal do cliente; primeiro sanear, depois operar.', 0)]))
A(bullet([('Pendências na conta fiscal', 1), (' — débitos em aberto ou situação cadastral irregular do estabelecimento.', 0)]))
A(bullet([('Crédito indisponível no mês', 1), (' — volume habilitado do detentor insuficiente para a demanda; por isso o dimensionamento antecipado.', 0)]))
A(bullet([('Considerar a economia antes da hora', 1), (' — a economia só é real após a homologação e a constatação da conta fiscal zerada; nunca antes.', 0)]))
A(question(23, 'Quando o produto NÃO é indicado?'))
A(bullet([('ICMS mensal baixo', 1), (' — o esforço operacional supera o ganho.', 0)]))
A(bullet([('Simples Nacional', 1), (' — regime de apuração distinto; fora do escopo do produto.', 0)]))
A(bullet([('Empresa com saldo credor próprio', 1), (' — ela não desembolsa ICMS, logo não há o que compensar. É perfil de potencial detentor/cedente — leve à área técnica, não venda o produto.', 0)]))
A(question(24, 'Qual a postura de compliance que o executivo deve transmitir?'))
A(bodyp([('A Lucro Ativo opera ', 0), ('créditos habilitados, por sistemas oficiais e sob homologação da SEFAZ/SP', 1),
         (', com registro integral (GIA, SPED, Conta Fiscal) e responsabilidade contratual de 5 anos. O executivo nunca promete percentual sem apuração: o que se garante é o rito, a comprovação e a condução técnica.', 0)]))

A(section('GLOSSÁRIO'))
for lead, rest in [
    ('Não-cumulatividade:', ' princípio (CF, art. 155, §2º, I) pelo qual se compensa o ICMS devido nas saídas com o ICMS cobrado nas entradas, evitando tributação "em cascata".'),
    ('Saldo credor:', ' excesso de créditos sobre débitos na apuração do período.'),
    ('Crédito acumulado:', ' parcela do saldo credor decorrente de hipótese legal específica (RICMS-SP, art. 71) que pode ser apropriada, utilizada e transferida.'),
    ('Transferência de crédito:', ' operação pela qual o detentor cede crédito acumulado a outro contribuinte, nas hipóteses do RICMS-SP.'),
    ('Cessionária:', ' a empresa que recebe o crédito transferido e o utiliza para compensar o próprio ICMS — o nosso cliente.'),
    ('Detentor (cedente):', ' quem acumulou o crédito e o cede com deságio; arca com as despesas do procedimento.'),
    ('Deságio:', ' desconto sobre o valor de face do crédito — é a fonte da economia do cliente.'),
    ('"Outros Créditos":', ' campo da GIA onde o crédito transferido é lançado (fundamento: art. 99 do RICMS-SP, conforme material-base).'),
    ('PRF:', ' formato do arquivo da GIA enviado pelo cliente para o lançamento; devolvemos o arquivo já transmitido.'),
    ('e-CredAc:', ' Sistema Eletrônico de Gerenciamento do Crédito Acumulado da SEFAZ-SP (Portaria SRE 65/2023) — a "conta eletrônica" dos créditos.'),
    ('GIA:', ' Guia de Informação e Apuração do ICMS; declaração mensal do contribuinte paulista.'),
    ('SPED Fiscal (EFD-ICMS/IPI):', ' escrituração fiscal digital; espelha a compensação e gera o recibo de entrega (comprovação).'),
    ('Conta Fiscal:', ' extrato/conta corrente fiscal do contribuinte na SEFAZ-SP; onde se constata guias quitadas e saldo zerado.'),
    ('DARE:', ' guia de recolhimento do saldo remanescente não compensado.'),
    ('Homologação:', ' ato pelo qual a SEFAZ-SP valida a compensação, autorizando seus efeitos.'),
    ('Glosa:', ' rejeição, total ou parcial, de crédito considerado indevido ou não comprovado.'),
]:
    A(bullet([(lead, 1), (rest, 0)]))

A(section('BASE LEGAL CITADA (para conferência do time técnico)'))
for lead, rest in [
    ('CF, art. 155, §2º', ' (não-cumulatividade do ICMS).'),
    ('LC 87/1996 (Lei Kandir)', ' — arts. 19-23 (direito ao crédito) e art. 25 (saldos credores e transferência de crédito acumulado).'),
    ('RICMS-SP (Decreto 45.490/2000)', ' — arts. 71 a 84 (crédito acumulado: hipóteses, apropriação, utilização e transferência).'),
    ('RICMS-SP, art. 99', ' — fundamento do lançamento em "Outros Créditos" na GIA/SPED (conforme material-base da operação).'),
    ('Portaria SRE 65/2023', ' — apropriação/utilização do crédito acumulado e sistema e-CredAc (fonte de dados: EFD desde 1º/jan/2026).'),
]:
    A(bullet([(lead, 1), (rest, 0)]))
A(box('atencao',
      'Todos os números comerciais (percentual de utilização, deságio/economia e condições) refletem o material-base vigente ',
      'e permanecem como PARÂMETRO A VALIDAR em proposta individual, até confirmação da área técnica/comercial da Lucro Ativo.'))
A(gap())

# ---------- APLICAR ----------
# 1) retitular o heading do capítulo
head = blocks[start]
runs = head.findall(qn('w:r'))
set_run_text(runs[0], 'Capítulo 2 · Compensação de ICMS com Créditos Acumulados — SP')
for r in runs[1:]:
    head.remove(r)

# 2) remover blocos antigos (start+1 .. end-1)
for el in blocks[start + 1:end]:
    body.remove(el)

# 3) inserir novos antes do Capítulo 3
cap3 = blocks[end]
for el in new_els:
    cap3.addprevious(el)

# 4) capa/intro: trocar "Recuperação de ICMS-SP" -> "Compensação de ICMS-SP"
for p in doc.paragraphs:
    if 'Recuperação de ICMS-SP' in p.text:
        for r in p.runs:
            if 'Recuperação de ICMS-SP' in r.text:
                r.text = r.text.replace('Recuperação de ICMS-SP', 'Compensação de ICMS-SP')

doc.save(OUT)
print('saved', OUT)

# verificação: nenhuma menção indevida restante
check = Document(OUT)
bad = []
intro_done = False
for p in check.paragraphs:
    t = p.text
    if 'Recuperação de ICMS' in t or 'recuperação de ICMS' in t:
        bad.append(t[:120])
print('menções "Recuperação de ICMS" restantes:', bad)
