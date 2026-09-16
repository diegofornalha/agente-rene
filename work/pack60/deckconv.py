#!/usr/bin/env python3
"""PDF de deck (Chrome/HTML, padrao Lucro Ativo) -> modelo geometrico.

Gera: (a) PPTX NATIVO EDITAVEL via python-pptx  (b) preview PNG via PIL
Ambos a partir do MESMO modelo, para que o preview mostre o que o PowerPoint mostrara.
"""
import json, os, statistics, sys

NAVY = 0x010860; GOLD = 0xC9A227; INK = 0x1A1F36; GREY = 0x5A6072
MONT = 'Montserrat'; INTER = 'Inter'
LOGO_WHITE = '/home/hermes/work/rota-fiscal/logo-white.png'
LOGO_NAVY  = '/home/hermes/work/rota-fiscal/logo-navy.png'


def _rgb(i):
    return (i >> 16 & 255, i >> 8 & 255, i & 255)


def _hex(t):
    return t[0] << 16 | t[1] << 8 | t[2]


def _fill_int(c):
    if c is None:
        return None
    return _hex(tuple(min(255, max(0, round(v * 255))) for v in c[:3]))


def _lum(i):
    r, g, b = _rgb(i)
    return 0.299 * r + 0.587 * g + 0.114 * b


def style_of(size_px, color, letterspaced, single_run, mixed_line):
    """Heuristica de fonte/peso calibrada no CSS real dos decks LA."""
    accent = color in (NAVY, GOLD) or _lum(color) > 200
    if letterspaced:
        return MONT, True
    if size_px >= 26:
        return MONT, True
    if size_px >= 17 and single_run and not mixed_line and accent:
        return MONT, True
    if color in (NAVY, GOLD):
        return INTER, True
    if _lum(color) > 200 and size_px >= 17 and single_run:
        return MONT, True
    return INTER, False


def extract_text(page, imgdir, pno):
    out = []
    raw = page.get_text('rawdict')
    for b in raw['blocks']:
        if b['type'] != 0:
            continue
        lines = []
        for l in b['lines']:
            chars = [c for s in l['spans'] for c in s['chars']]
            if not chars:
                continue
            txt_all = ''.join(c['c'] for c in chars)
            if not txt_all.strip():
                continue
            nsp = sum(1 for c in chars if c['c'] == ' ')
            solid = [c for c in chars if c['c'] != ' ']
            sizes = [s['size'] for s in l['spans']]
            size = max(sizes)
            ls = False; spc = 0
            if len(chars) > 4 and nsp / len(chars) > 0.34:
                gaps = [round(b2['bbox'][0] - a['bbox'][2], 2) for a, b2 in zip(solid, solid[1:])]
                inner = [g for g in gaps if g < 0.20 * size]
                if inner:
                    ls = True
                    spc = max(0.0, statistics.median(inner))
            runs = []
            if ls:
                thr = 0.20 * size
                buf = ''
                prev = None
                for c in solid:
                    if prev is not None and (c['bbox'][0] - prev['bbox'][2]) > thr:
                        buf += ' '
                    buf += c['c']
                    prev = c
                col = l['spans'][0]['color']
                runs = [{'t': buf, 'size': size, 'color': col}]
            else:
                for s in l['spans']:
                    t = ''.join(c['c'] for c in s['chars'])
                    if not t:
                        continue
                    if runs and runs[-1]['color'] == s['color'] and abs(runs[-1]['size'] - s['size']) < .2:
                        runs[-1]['t'] += t
                    else:
                        runs.append({'t': t, 'size': s['size'], 'color': s['color']})
            runs = [r for r in runs if r['t'] != '']
            if not runs:
                continue
            mixed = len({r['color'] for r in runs}) > 1
            for r in runs:
                px = r['size'] / 0.75
                fam, bold = style_of(px, r['color'], ls, len(runs) == 1, mixed)
                r['font'] = fam; r['bold'] = bold
                r['spc'] = round(spc, 2) if ls else 0
            lines.append({'bbox': [round(v, 2) for v in l['bbox']], 'runs': runs,
                          'size': round(size, 2), 'ls': ls,
                          'asc': round(l['spans'][0].get('ascender', .9), 3),
                          'desc': round(l['spans'][0].get('descender', -.2), 3)})
        if not lines:
            continue
        out.append({'kind': 'text', 'bbox': [round(v, 2) for v in b['bbox']], 'lines': lines})
    return out


def extract_images(page, imgdir, pno):
    res = []
    d = page.get_text('dict')
    k = 0
    for b in d['blocks']:
        if b['type'] != 1:
            continue
        data = b.get('image')
        if not data:
            continue
        ext = b.get('ext', 'png')
        fn = os.path.join(imgdir, 'p%02d_i%02d.%s' % (pno, k, ext))
        with open(fn, 'wb') as f:
            f.write(data)
        res.append({'kind': 'img', 'bbox': [round(v, 2) for v in b['bbox']], 'file': fn,
                    'w': b.get('width'), 'h': b.get('height')})
        k += 1
    return res


def extract_shapes(page):
    rects, paths = [], []
    for x in page.get_drawings():
        r = x['rect']
        w, h = r.x1 - r.x0, r.y1 - r.y0
        if w <= 0.4 or h <= 0.4:
            if w <= 0 or h <= 0:
                continue
        has_re = any(it[0] == 're' for it in x['items'])
        curvy = any(it[0] == 'c' for it in x['items'])
        nseg = len(x['items'])
        if not has_re and curvy and nseg > 2 and w < 95 and h < 95:
            paths.append({'bbox': [r.x0, r.y0, r.x1, r.y1], 'fill': _fill_int(x['fill'])})
            continue
        rad = 0.0
        if curvy and not has_re:
            it0 = x['items'][0]
            try:
                rad = max(0.0, min(it0[1].x, it0[2].x) - r.x0)
            except Exception:
                rad = 7.5
            if rad <= 0.2:
                rad = 7.5
        rects.append({'kind': 'rect', 'bbox': [round(v, 2) for v in (r.x0, r.y0, r.x1, r.y1)],
                      'fill': _fill_int(x['fill']), 'stroke': _fill_int(x['color']),
                      'lw': round(x.get('width') or 0.75, 2), 'rad': round(rad, 2)})
    return rects, paths


def cluster_logo(paths):
    """Aglomera paths vetoriais pequenos (logo SVG) num unico grupo."""
    if len(paths) < 8:
        return [], paths
    xs = [p['bbox'] for p in paths]
    x0 = min(b[0] for b in xs); y0 = min(b[1] for b in xs)
    x1 = max(b[2] for b in xs); y1 = max(b[3] for b in xs)
    if (x1 - x0) > 620 or (y1 - y0) > 260:
        return [], paths
    fills = [p['fill'] for p in paths if p['fill'] is not None]
    light = fills and statistics.mean(_lum(f) for f in fills) > 128
    return [{'kind': 'logo', 'bbox': [round(x0, 2), round(y0, 2), round(x1, 2), round(y1, 2)],
             'file': LOGO_WHITE if light else LOGO_NAVY}], []


def build_model(pdf, outdir):
    import pymupdf
    imgdir = os.path.join(outdir, 'img'); os.makedirs(imgdir, exist_ok=True)
    doc = pymupdf.open(pdf)
    W, H = doc[0].rect.x1, doc[0].rect.y1
    slides = []
    for pno, page in enumerate(doc, 1):
        imgs = extract_images(page, imgdir, pno)
        rects, paths = extract_shapes(page)
        logo, leftover = cluster_logo(paths)
        area = W * H
        bg = [i for i in imgs if (i['bbox'][2] - i['bbox'][0]) * (i['bbox'][3] - i['bbox'][1]) > .8 * area]
        fg = [i for i in imgs if i not in bg]
        txt = extract_text(page, imgdir, pno)
        shapes = bg + rects + fg + logo + txt
        for p in leftover:
            shapes.append({'kind': 'rect', 'bbox': [round(v, 2) for v in p['bbox']],
                           'fill': p['fill'], 'stroke': None, 'lw': 0, 'rad': 0})
        slides.append({'n': pno, 'shapes': shapes})
    return {'w': W, 'h': H, 'slides': slides, 'src': pdf}


# ─────────────────────────── PPTX ───────────────────────────
def to_pptx(model, out, title=None):
    from pptx import Presentation
    from pptx.util import Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    E = 12700  # EMU por ponto

    prs = Presentation()
    prs.slide_width = Emu(int(model['w'] * E)); prs.slide_height = Emu(int(model['h'] * E))
    blank = prs.slide_layouts[6]

    for sl in model['slides']:
        s = prs.slides.add_slide(blank)
        for sh in sl['shapes']:
            x0, y0, x1, y1 = sh['bbox']
            w, h = max(1.0, x1 - x0), max(1.0, y1 - y0)
            if sh['kind'] in ('img', 'logo'):
                s.shapes.add_picture(sh['file'], Emu(int(x0 * E)), Emu(int(y0 * E)),
                                     Emu(int(w * E)), Emu(int(h * E)))
            elif sh['kind'] == 'rect':
                shp_t = MSO_SHAPE.ROUNDED_RECTANGLE if sh['rad'] > 0.5 else MSO_SHAPE.RECTANGLE
                shp = s.shapes.add_shape(shp_t, Emu(int(x0 * E)), Emu(int(y0 * E)),
                                         Emu(int(w * E)), Emu(int(h * E)))
                if shp_t == MSO_SHAPE.ROUNDED_RECTANGLE:
                    try:
                        shp.adjustments[0] = min(0.5, sh['rad'] / min(w, h))
                    except Exception:
                        pass
                if sh['fill'] is None:
                    shp.fill.background()
                else:
                    shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor(*_rgb(sh['fill']))
                if sh['stroke'] is None:
                    shp.line.fill.background()
                else:
                    shp.line.color.rgb = RGBColor(*_rgb(sh['stroke']))
                    shp.line.width = Pt(max(0.5, sh['lw']))
                shp.shadow.inherit = False
                shp.text_frame.word_wrap = True
            else:  # text
                lines = sh['lines']
                first = lines[0]
                sz0 = first['size']
                top = first['bbox'][1] - 0.06 * sz0
                # largura: maior linha + folga (evita re-wrap; layout = PDF)
                maxw = max(l['bbox'][2] - l['bbox'][0] for l in lines)
                left = min(l['bbox'][0] for l in lines)
                cx = [(l['bbox'][0] + l['bbox'][2]) / 2 for l in lines]
                centered = len(lines) > 1 and (max(cx) - min(cx)) < 2.0 and \
                    (max(l['bbox'][0] for l in lines) - left) > 3.0
                slack = maxw * 0.10 + 6
                if centered:
                    bx = min(cx) - (maxw + slack) / 2
                else:
                    bx = left - 1.0
                bw = maxw + slack
                bh = (lines[-1]['bbox'][3] - top) + 0.35 * lines[-1]['size'] + 2
                tb = s.shapes.add_textbox(Emu(int(bx * E)), Emu(int(top * E)),
                                          Emu(int(bw * E)), Emu(int(bh * E)))
                tf = tb.text_frame; tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                prev_base = None
                for i, l in enumerate(lines):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
                    base = l['bbox'][1]
                    if prev_base is not None:
                        pitch = base - prev_base
                        p.line_spacing = max(0.75, min(2.4, pitch / l['size']))
                    else:
                        p.line_spacing = 1.0
                    prev_base = base
                    for r in l['runs']:
                        run = p.add_run(); run.text = r['t']
                        f = run.font
                        f.name = r['font']; f.size = Pt(round(r['size'], 1)); f.bold = bool(r['bold'])
                        f.color.rgb = RGBColor(*_rgb(r['color']))
                        if r.get('spc'):
                            run._r.get_or_add_rPr().set('spc', str(int(round(r['spc'] * 100))))
    if title:
        prs.core_properties.title = title
    prs.save(out)
    return out


# ─────────────────────────── PREVIEW ───────────────────────────
FONTS = {('Montserrat', True): '/home/hermes/.local/share/fonts/Montserrat-Bold.ttf',
         ('Montserrat', False): '/home/hermes/.local/share/fonts/Montserrat-Bold.ttf',
         ('Inter', True): '/home/hermes/.local/share/fonts/Inter-Regular.ttf',
         ('Inter', False): '/home/hermes/.local/share/fonts/Inter-Regular.ttf'}
_fc = {}


def _font(fam, bold, px):
    from PIL import ImageFont
    k = (fam, bold, int(px))
    if k not in _fc:
        _fc[k] = ImageFont.truetype(FONTS[(fam, bold)], max(6, int(px)))
    return _fc[k]


def to_preview(model, outdir, scale=1.5):
    from PIL import Image, ImageDraw
    os.makedirs(outdir, exist_ok=True)
    files = []
    W, H = int(model['w'] * scale), int(model['h'] * scale)
    for sl in model['slides']:
        im = Image.new('RGB', (W, H), 'white')
        dr = ImageDraw.Draw(im, 'RGBA')
        for sh in sl['shapes']:
            x0, y0, x1, y1 = [v * scale for v in sh['bbox']]
            if sh['kind'] in ('img', 'logo'):
                try:
                    src = Image.open(sh['file']).convert('RGBA')
                    src = src.resize((max(1, int(x1 - x0)), max(1, int(y1 - y0))), Image.LANCZOS)
                    im.paste(src, (int(x0), int(y0)), src)
                except Exception:
                    pass
            elif sh['kind'] == 'rect':
                fill = _rgb(sh['fill']) if sh['fill'] is not None else None
                out = _rgb(sh['stroke']) if sh['stroke'] is not None else None
                r = sh['rad'] * scale
                if r > 1:
                    dr.rounded_rectangle([x0, y0, x1, y1], radius=min(r, (min(x1 - x0, y1 - y0)) / 2),
                                         fill=fill, outline=out, width=max(1, int(sh['lw'] * scale)))
                else:
                    dr.rectangle([x0, y0, x1, y1], fill=fill, outline=out,
                                 width=max(1, int((sh['lw'] or 1) * scale)))
            else:
                for l in sh['lines']:
                    bx = l['bbox'][0] * scale
                    base_top = l['bbox'][1] * scale
                    x = bx
                    for r in l['runs']:
                        f = _font(r['font'], r['bold'], r['size'] * scale)
                        sp = (r.get('spc') or 0) * scale
                        col = _rgb(r['color'])
                        sw = 0.6 if (r['bold'] and r['font'] == 'Inter') else 0
                        if sp:
                            for ch in r['t']:
                                dr.text((x, base_top), ch, font=f, fill=col,
                                        stroke_width=sw, stroke_fill=col)
                                x += dr.textlength(ch, font=f) + sp
                        else:
                            dr.text((x, base_top), r['t'], font=f, fill=col,
                                    stroke_width=sw, stroke_fill=col)
                            x += dr.textlength(r['t'], font=f)
        fn = os.path.join(outdir, 'prev_%02d.png' % sl['n'])
        im.save(fn); files.append(fn)
    return files


if __name__ == '__main__':
    pdf, outdir, pptx = sys.argv[1], sys.argv[2], sys.argv[3]
    os.makedirs(outdir, exist_ok=True)
    m = build_model(pdf, outdir)
    json.dump(m, open(os.path.join(outdir, 'model.json'), 'w'))
    to_pptx(m, pptx, title=os.path.basename(pptx))
    n = sum(len(s['shapes']) for s in m['slides'])
    print('slides=%d shapes=%d -> %s' % (len(m['slides']), n, pptx))
