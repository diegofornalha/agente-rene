#!/usr/bin/env python3
"""Deck PDF (padrao Lucro Ativo, gerado do HTML) -> PPTX NATIVO com texto editavel.

Estrategia hibrida (fidelidade total + edicao real):
  camada 1  fundo = a propria pagina do PDF renderizada SEM o texto (redaction so de
            texto; imagens e vetores preservados) -> design identico, pixel a pixel
  camada 2  texto = caixas de texto nativas do PowerPoint, com fonte/tamanho/cor/peso
            e posicao extraidos do PDF -> tudo editavel

Uso: deckconv3.py entrada.pdf saida.pptx [outdir]
"""
import os, sys, json

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deckconv2 import extract_text, _rgb  # noqa: E402

BG_SCALE = 2.0


def build_model(pdf, outdir):
    import pymupdf
    os.makedirs(outdir, exist_ok=True)
    imgdir = os.path.join(outdir, 'img'); os.makedirs(imgdir, exist_ok=True)
    doc = pymupdf.open(pdf)
    W, H = doc[0].rect.x1, doc[0].rect.y1
    slides = []
    for pno, page in enumerate(doc, 1):
        blocks = extract_text(page, imgdir, pno)          # texto ANTES de apagar
        for b in blocks:                                   # marca cada linha
            for l in b['lines']:
                r = pymupdf.Rect(l['bbox'])
                r.x0 -= 0.6; r.y0 -= 0.6; r.x1 += 0.6; r.y1 += 0.6
                page.add_redact_annot(r)
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                              graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                              text=pymupdf.PDF_REDACT_TEXT_REMOVE)
        pm = page.get_pixmap(matrix=pymupdf.Matrix(BG_SCALE, BG_SCALE), alpha=False)
        bg = os.path.join(imgdir, 'bg_%02d.png' % pno)
        pm.save(bg)
        left = len(page.get_text().strip())
        slides.append({'n': pno, 'bg': bg, 'text': blocks, 'leftover': left})
    return {'w': W, 'h': H, 'slides': slides, 'src': pdf}


def _geom(lines):
    """caixa que envolve o bloco + deteccao de centralizacao (igual ao PDF)."""
    first = lines[0]
    top = first['bbox'][1] - 0.06 * first['size']
    maxw = max(l['bbox'][2] - l['bbox'][0] for l in lines)
    left = min(l['bbox'][0] for l in lines)
    cx = [(l['bbox'][0] + l['bbox'][2]) / 2 for l in lines]
    centered = len(lines) > 1 and (max(cx) - min(cx)) < 2.0 and \
        (max(l['bbox'][0] for l in lines) - left) > 3.0
    slack = maxw * 0.10 + 6
    bx = (min(cx) - (maxw + slack) / 2) if centered else (left - 1.0)
    bw = maxw + slack
    bh = (lines[-1]['bbox'][3] - top) + 0.35 * lines[-1]['size'] + 2
    return bx, top, bw, bh, centered


def to_pptx(model, out, title=None):
    from pptx import Presentation
    from pptx.util import Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    E = 12700
    prs = Presentation()
    prs.slide_width = Emu(int(model['w'] * E)); prs.slide_height = Emu(int(model['h'] * E))
    blank = prs.slide_layouts[6]
    for sl in model['slides']:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(sl['bg'], 0, 0,
                             Emu(int(model['w'] * E)), Emu(int(model['h'] * E)))
        for b in sl['text']:
            lines = b['lines']
            bx, top, bw, bh, centered = _geom(lines)
            tb = s.shapes.add_textbox(Emu(int(bx * E)), Emu(int(top * E)),
                                      Emu(int(bw * E)), Emu(int(bh * E)))
            tf = tb.text_frame; tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.TOP
            prev = None
            for i, l in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
                base = l['bbox'][1]
                p.line_spacing = 1.0 if prev is None else \
                    max(0.75, min(2.4, (base - prev) / l['size']))
                prev = base
                for r in l['runs']:
                    run = p.add_run(); run.text = r['t']
                    f = run.font
                    f.name = r['font']; f.size = Pt(round(r['size'], 1))
                    f.bold = bool(r['bold'])
                    f.color.rgb = RGBColor(*_rgb(r['color']))
                    if r.get('spc'):
                        run._r.get_or_add_rPr().set('spc', str(int(round(r['spc'] * 100))))
    if title:
        prs.core_properties.title = title
    prs.save(out)
    return out


# ---------------------------------------------------------------- preview
_VAR = {('Montserrat', True): ('/home/hermes/.local/share/fonts/Montserrat-Bold.ttf', 'Bold'),
        ('Montserrat', False): ('/home/hermes/.local/share/fonts/Montserrat-Bold.ttf', 'Regular'),
        ('Inter', True): ('/home/hermes/.local/share/fonts/Inter-Regular.ttf', 'Bold'),
        ('Inter', False): ('/home/hermes/.local/share/fonts/Inter-Regular.ttf', 'Regular')}
_cache = {}


def _font(fam, bold, px):
    from PIL import ImageFont
    k = (fam, bold, int(px))
    if k not in _cache:
        path, style = _VAR[(fam, bold)]
        f = ImageFont.truetype(path, max(6, int(px)))
        try:
            f.set_variation_by_name(style)
        except Exception:
            pass
        _cache[k] = f
    return _cache[k]


def to_preview(model, outdir, scale=1.5):
    """Composicao fundo + texto (aproxima o que o PowerPoint vai mostrar)."""
    from PIL import Image, ImageDraw
    os.makedirs(outdir, exist_ok=True)
    files = []
    W, H = int(model['w'] * scale), int(model['h'] * scale)
    for sl in model['slides']:
        im = Image.open(sl['bg']).convert('RGB').resize((W, H), Image.LANCZOS)
        dr = ImageDraw.Draw(im)
        for b in sl['text']:
            for l in b['lines']:
                x = l['bbox'][0] * scale
                y = l['bbox'][1] * scale
                for r in l['runs']:
                    f = _font(r['font'], r['bold'], r['size'] * scale)
                    col = _rgb(r['color'])
                    sp = (r.get('spc') or 0) * scale
                    if sp:
                        for ch in r['t']:
                            dr.text((x, y), ch, font=f, fill=col)
                            x += dr.textlength(ch, font=f) + sp
                    else:
                        dr.text((x, y), r['t'], font=f, fill=col)
                        x += dr.textlength(r['t'], font=f)
        fn = os.path.join(outdir, 'prev_%02d.png' % sl['n'])
        im.save(fn); files.append(fn)
    return files


if __name__ == '__main__':
    pdf, pptx = sys.argv[1], sys.argv[2]
    outdir = sys.argv[3] if len(sys.argv) > 3 else '/tmp/dc3_' + os.path.basename(pptx)[:12]
    m = build_model(pdf, outdir)
    json.dump({k: v for k, v in m.items() if k != 'slides'} | {'n': len(m['slides'])},
              open(os.path.join(outdir, 'meta.json'), 'w'))
    to_pptx(m, pptx, title=os.path.splitext(os.path.basename(pptx))[0])
    nt = sum(sum(len(b['lines']) for b in s['text']) for s in m['slides'])
    left = sum(s['leftover'] for s in m['slides'])
    print('slides=%d linhas_texto=%d texto_residual_no_fundo=%d -> %s'
          % (len(m['slides']), nt, left, pptx))
    json.dump(m, open(os.path.join(outdir, 'model.json'), 'w'))
