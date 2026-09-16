# -*- coding: utf-8 -*-
"""Cronograma Semanal de Gestão — Parcerias (Ulisses) — gerador do XLSX.
Fonte: Playbook Gestão de Parcerias v1.0 (minuta). Volumes = PARÂMETRO A VALIDAR."""
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

NAVY = "010860"
NAVY_FILL = PatternFill("solid", fgColor=NAVY)
WHITE_BOLD = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, color=NAVY, size=16)
SUB_FONT = Font(name="Calibri", color="555566", size=10, italic=True)
BODY = Font(name="Calibri", size=10, color="1A1A2E")
BODY_B = Font(name="Calibri", size=10, bold=True, color=NAVY)
ZEBRA = PatternFill("solid", fgColor="F4F5FA")
THIN = Side(style="thin", color="C9CCDD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
WRAP_C = Alignment(wrap_text=True, vertical="center", horizontal="center")
LEFT_C = Alignment(wrap_text=True, vertical="center", horizontal="left")

# categorias e cores
CAT = {
    "CART": ("Gestão de carteira", "D9E0F5"),
    "PROS": ("Prospecção de parceiros", "F3E6C8"),
    "PLAN": ("Planejamento conjunto", "DCEDE2"),
    "CAMPO": ("Campo & visitas", "E7E0F1"),
    "ANAL": ("Análise & rituais internos", "E9EAEF"),
}
CAT_FILL = {k: PatternFill("solid", fgColor=v[1]) for k, v in CAT.items()}


def sheet_title(ws, row, title, subtitle=None):
    ws.cell(row=row, column=1, value=title).font = TITLE_FONT
    if subtitle:
        ws.cell(row=row + 1, column=1, value=subtitle).font = SUB_FONT
        return row + 3
    return row + 2


def header_row(ws, row, headers, widths=None, start_col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.fill = NAVY_FILL
        c.font = WHITE_BOLD
        c.alignment = WRAP_C
        c.border = BORDER
        if widths:
            ws.column_dimensions[get_column_letter(start_col + i)].width = widths[i]
    ws.row_dimensions[row].height = 26


wb = Workbook()

# ------------------------------------------------------------------ ABA 1
ws = wb.active
ws.title = "Semana de Gestão"
ws.sheet_view.showGridLines = False
r = sheet_title(
    ws, 1, "Semana-padrão de gestão — Coordenação de Parcerias",
    "Rotina temática do coordenador (planejar → executar → registrar → analisar) · régua por tier: A quinzenal · B mensal · C automatizada · Playbook v1 — PARÂMETRO A VALIDAR",
)
DAYS = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta"]
BLOCKS = ["08h30–10h", "10h–12h", "13h30–15h30", "15h30–17h30", "17h30–18h"]
# grade: (texto, categoria)
GRID = [
    [  # 08h30–10h
        ("Planejamento da semana — leitura do dashboard (Twenty) e seleção dos parceiros do ciclo", "ANAL"),
        ("Preparação da Weekly de Parcerias — pipeline de indicados, onboarding e exceções", "ANAL"),
        ("Contatos de gestão — Tier A (ciclo quinzenal): pipeline da carteira e próxima abertura", "CART"),
        ("Preparação dos planos conjuntos — contas-alvo × ICP (5–10 contas por ciclo)", "PLAN"),
        ("Contatos de gestão — recuperação da régua da semana (aderência ≥ 90%)", "CART"),
    ],
    [  # 10h–12h
        ("Contatos de gestão — Tier A (ciclo quinzenal)", "CART"),
        ("Weekly de Parcerias com o CCO (terças · 45 min) + desdobramentos no CRM", "ANAL"),
        ("Visita presencial / agenda de campo (1–2 por semana)", "CAMPO"),
        ("Reuniões de planejamento & resultado com parceiros (Tier A mensal · Tier B bimestral)", "PLAN"),
        ("Reuniões de apresentação — proposta de trabalho a parceiros-alvo", "PROS"),
    ],
    [  # 13h30–15h30
        ("Contatos de gestão — Tier A / Tier B (mensal)", "CART"),
        ("Reuniões de apresentação — proposta de trabalho (modelo, papéis, régua, dashboard)", "PROS"),
        ("Visita presencial / reunião de apresentação em campo", "CAMPO"),
        ("Reuniões de planejamento & resultado com parceiros", "PLAN"),
        ("Análise da semana — aderência à régua, Funil 1 (credenciamentos) e Funil 2 (indicações)", "ANAL"),
    ],
    [  # 15h30–17h30
        ("Prospecção de novos parceiros — entidades (SESCON, ACIs), indicações e follow-ups", "PROS"),
        ("Contatos de gestão — Tier B (ciclo mensal): status das indicações e destrave de contas", "CART"),
        ("Prospecção de novos parceiros — follow-up de propostas e credenciamentos em curso", "PROS"),
        ("Contatos de gestão — Tier A/B + reativação de inativos (campanha de conta-alvo)", "CART"),
        ("Registro consolidado + preparação da próxima semana e da pauta da Weekly", "ANAL"),
    ],
    [  # 17h30–18h
        ("Registro do dia no CRM — atividades e próximos passos datados", "ANAL"),
        ("Registro do dia no CRM — atividades e próximos passos datados", "ANAL"),
        ("Registro do dia no CRM — atividades e próximos passos datados", "ANAL"),
        ("Registro do dia no CRM — atividades e próximos passos datados", "ANAL"),
        ("Fechamento — leitura de gestão da semana (o que a carteira está dizendo)", "ANAL"),
    ],
]
ws.column_dimensions["A"].width = 13
for i in range(5):
    ws.column_dimensions[get_column_letter(2 + i)].width = 34
header_row(ws, r, ["Bloco"] + DAYS)
for bi, block in enumerate(BLOCKS):
    row = r + 1 + bi
    c = ws.cell(row=row, column=1, value=block)
    c.fill = NAVY_FILL
    c.font = WHITE_BOLD
    c.alignment = WRAP_C
    c.border = BORDER
    ws.row_dimensions[row].height = 64 if bi < 4 else 40
    for di in range(5):
        text, cat = GRID[bi][di]
        cell = ws.cell(row=row, column=2 + di, value=text)
        cell.fill = CAT_FILL[cat]
        cell.font = Font(name="Calibri", size=9.5, color="1A1A2E")
        cell.alignment = WRAP
        cell.border = BORDER
# legenda
lr = r + len(BLOCKS) + 2
ws.cell(row=lr, column=1, value="Legenda").font = BODY_B
for i, (k, (name, color)) in enumerate(CAT.items()):
    c = ws.cell(row=lr + 1 + i, column=1)
    c.fill = PatternFill("solid", fgColor=color)
    c.border = BORDER
    ws.cell(row=lr + 1 + i, column=2, value=name).font = BODY
nr = lr + len(CAT) + 2
ws.cell(row=nr, column=1, value="Volumes de referência da semana (PARÂMETRO A VALIDAR): 20–25 ligações de gestão · 10–12 ligações de prospecção · 3–4 reuniões de apresentação · 5–6 reuniões de planejamento & resultado · 1–2 visitas presenciais · evento a cada 45–60 dias.").font = SUB_FONT
ws.merge_cells(start_row=nr, start_column=1, end_row=nr, end_column=6)
ws.row_dimensions[nr].height = 26
ws.cell(row=nr, column=1).alignment = WRAP
ws.freeze_panes = "B5"

# ------------------------------------------------------------------ ABA 2
ws = wb.create_sheet("Planner Semanal")
ws.sheet_view.showGridLines = False
r = sheet_title(
    ws, 1, "Planner semanal do coordenador — modelo replicável",
    "Duplicar esta aba a cada semana (Semana de ___/___ a ___/___). Instrumento de planejamento e leitura — não de cobrança.",
)
heads = ["Dia", "Foco de gestão do dia", "Parceiros do ciclo (tier)",
         "Reuniões agendadas", "Registros no CRM", "Observações do coordenador"]
widths = [12, 32, 28, 28, 26, 32]
header_row(ws, r, heads, widths)
days = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta"]
for i, d in enumerate(days):
    row = r + 1 + i
    ws.row_dimensions[row].height = 52
    c = ws.cell(row=row, column=1, value=d)
    c.font = BODY_B
    c.alignment = LEFT_C
    c.border = BORDER
    if i % 2 == 1:
        c.fill = ZEBRA
    for j in range(2, 7):
        cell = ws.cell(row=row, column=j, value="")
        cell.border = BORDER
        cell.alignment = WRAP
        cell.font = BODY
        if i % 2 == 1:
            cell.fill = ZEBRA
frow = r + 6
ws.row_dimensions[frow].height = 56
c = ws.cell(row=frow, column=1, value="Fechamento da semana")
c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
c.fill = NAVY_FILL
c.alignment = LEFT_C
c.border = BORDER
labels = ["O que a semana ensinou (leitura de gestão)", "Parceiros que pedem atenção no próximo ciclo",
          "Pauta levada à Weekly com o CCO", "Pendências de registro", "Ajustes na semana seguinte"]
for j, lab in enumerate(labels):
    cell = ws.cell(row=frow, column=2 + j, value=lab + ":")
    cell.font = Font(name="Calibri", size=9, italic=True, color="555566")
    cell.border = BORDER
    cell.alignment = WRAP
ws.freeze_panes = "B4"

# ------------------------------------------------------------------ ABA 3
ws = wb.create_sheet("Cadência da Carteira")
r = sheet_title(
    ws, 1, "Cadência da carteira — régua de contato por tier",
    "Tier A = contato efetivo quinzenal (14 dias) · Tier B = mensal (30 dias) · Tier C = régua automatizada (Agente + MKT). Contato efetivo = conversa real, não mensagem sem resposta.",
)
heads = ["Parceiro", "Tier", "Último contato", "Próximo contato", "Plano conjunto",
         "Status do ciclo", "Notas do coordenador"]
widths = [34, 8, 15, 16, 14, 18, 42]
header_row(ws, r, heads, widths)
first_data = r + 1
examples = [
    ("Contabilidade Exemplo A (fictício)", "A", date(2026, 8, 31), "Sim", "Em dia", "Plano conjunto do ciclo vigente; 2 contas-alvo em abertura."),
    ("BPO Financeiro Exemplo (fictício)", "A", date(2026, 8, 17), "Sim", "Contato pendente", "Reagendar contato quinzenal; conferir indicação no funil."),
    ("Consultoria de Negócios Exemplo (fictício)", "B", date(2026, 8, 20), "Sim", "Em dia", "Reunião de resultado bimestral prevista para o próximo ciclo."),
    ("Escritório Jurídico Exemplo (fictício)", "B", date(2026, 8, 10), "Não", "Reativação", "Perfil jurídico: credenciamento condicionado ao parecer do CLO."),
    ("Representante Comercial Exemplo (fictício)", "C", date(2026, 7, 15), "Não", "Régua automatizada", "Tier C: conteúdo, convites a eventos e gatilhos via Agente/MKT."),
]
total_rows = 50  # 5 exemplos + 45 em branco
for i in range(total_rows):
    row = first_data + i
    ws.row_dimensions[row].height = 20
    for j in range(1, 8):
        cell = ws.cell(row=row, column=j)
        cell.border = BORDER
        cell.font = BODY
        cell.alignment = LEFT_C if j != 2 else WRAP_C
        if i % 2 == 1:
            cell.fill = ZEBRA
    if i < len(examples):
        nome, tier, ultimo, plano, status, nota = examples[i]
        ws.cell(row=row, column=1, value=nome)
        ws.cell(row=row, column=2, value=tier)
        c = ws.cell(row=row, column=3, value=ultimo)
        c.number_format = "DD/MM/YYYY"
        ws.cell(row=row, column=5, value=plano)
        ws.cell(row=row, column=6, value=status)
        ws.cell(row=row, column=7, value=nota)
    else:
        ws.cell(row=row, column=3).number_format = "DD/MM/YYYY"
    # fórmula do próximo contato (A=+14, B=+30, C=régua automatizada)
    f = ('=IF(OR($B{0}="",$C{0}=""),"",IF($B{0}="A",$C{0}+14,'
         'IF($B{0}="B",$C{0}+30,"Régua automatizada")))').format(row)
    fc = ws.cell(row=row, column=4, value=f)
    fc.number_format = "DD/MM/YYYY"
last_data = first_data + total_rows - 1
# validações
dv_tier = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True,
                         promptTitle="Tier", prompt="A = quinzenal · B = mensal · C = automatizada")
dv_plano = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
dv_status = DataValidation(type="list",
                           formula1='"Em dia,Contato pendente,Reativação,Stand-by,Régua automatizada"',
                           allow_blank=True)
ws.add_data_validation(dv_tier); ws.add_data_validation(dv_plano); ws.add_data_validation(dv_status)
dv_tier.add(f"B{first_data}:B{last_data}")
dv_plano.add(f"E{first_data}:E{last_data}")
dv_status.add(f"F{first_data}:F{last_data}")
# formatação condicional no próximo contato
rng = f"D{first_data}:D{last_data}"
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND(ISNUMBER($D{first_data}),$D{first_data}<TODAY())'],
    fill=PatternFill("solid", fgColor="F6D3D3"),
    font=Font(color="8B1A1A", bold=True), stopIfTrue=True))
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND(ISNUMBER($D{first_data}),$D{first_data}<=TODAY()+3)'],
    fill=PatternFill("solid", fgColor="FBEFC9"),
    font=Font(color="7A5A00"), stopIfTrue=True))
ws.freeze_panes = f"A{first_data}"

# ------------------------------------------------------------------ ABA 4
ws = wb.create_sheet("Rituais & Ciclo")
ws.sheet_view.showGridLines = False
r = sheet_title(
    ws, 1, "Rituais & ciclo de gestão do canal",
    "Calendário de gestão do Playbook v1 (Seção 6.7) — cada ritual declara o que o alimenta (entrada) e o que produz (saída).",
)
heads = ["Ritual", "Frequência", "Propósito de gestão", "Participantes",
         "Entrada (o que alimenta)", "Saída (o que produz)"]
widths = [24, 16, 34, 24, 30, 30]
header_row(ws, r, heads, widths)
rituais = [
    ("Weekly de Parcerias (CCO)", "Semanal · terças · 45 min",
     "Alinhamento do canal: mapeamento, pipeline de indicados, onboarding e exceções",
     "Ulisses + Cleber Candia (CCO & CSO)",
     "Planner da semana, dashboard do funil, pendências de roteamento",
     "Decisões de exceção, prioridades da semana, pauta registrada"),
    ("Contato efetivo Tier A", "Quinzenal (a cada 14 dias)",
     "Manter a carteira estratégica em movimento: pipeline, contas-alvo e próxima abertura",
     "Coordenador + parceiro",
     "Cadência da carteira, plano conjunto vigente, status das indicações",
     "Atividade no CRM + próximo passo datado"),
    ("Contato efetivo Tier B", "Mensal (a cada 30 dias)",
     "Acompanhar status das indicações e destravar contas",
     "Coordenador + parceiro",
     "Cadência da carteira, funil de indicações",
     "Atividade no CRM + próximo passo datado"),
    ("Régua Tier C (automatizada)", "Contínua",
     "Manter presença de marca e gatilhos de reativação sem consumir agenda",
     "Agente de CRM + Marketing",
     "Base Tier C segmentada, calendário de conteúdo/eventos",
     "Disparos registrados, alertas de inatividade, gatilhos de reativação"),
    ("Reunião de planejamento & resultado", "Mensal (Tier A) · Bimestral (Tier B)",
     "Planejamento & resultado com o parceiro: conversões, qualidade da indicação e plano do próximo ciclo",
     "Coordenador + parceiro (+ executivo se houver conta ativa)",
     "Plano conjunto, pipeline da carteira, feedback de vendas",
     "Ata + plano conjunto revisado + saúde das contas (sensor de CS)"),
    ("Revisão mensal de carteira", "Mensal",
     "Leitura de gestão da base: tiers, status, aderência à régua e reativações",
     "Coordenador (prepara) + CCO (valida)",
     "Cadência da carteira, Painel do Coordenador do mês",
     "Reclassificação de tier/status, campanhas de conta-alvo, pauta da Weekly"),
    ("QBR do canal (parceiros A)", "Trimestral",
     "Resultado do trimestre, ranking e metas do próximo ciclo com os top parceiros",
     "CCO + Coordenador + top parceiros",
     "Painel trimestral, ranking de indicações convertidas",
     "Metas do próximo ciclo, reconhecimento, ajustes de plano"),
    ("Evento / palestra do canal", "A cada 45–60 dias",
     "Autoridade técnica (Reforma Tributária etc.): capta parceiro novo e reativa a base no mesmo movimento",
     "Rede de parceiros + convidados + LA",
     "Calendário de MKT, lista de inativos/alvos, tema técnico",
     "Lista de presença + follow-up em 5 dias úteis + reativações"),
]
for i, row_data in enumerate(rituais):
    row = r + 1 + i
    ws.row_dimensions[row].height = 46
    for j, v in enumerate(row_data):
        cell = ws.cell(row=row, column=1 + j, value=v)
        cell.border = BORDER
        cell.alignment = WRAP
        cell.font = BODY_B if j == 0 else BODY
        if i % 2 == 1:
            cell.fill = ZEBRA
ws.freeze_panes = f"A{r+1}"

# ------------------------------------------------------------------ ABA 5
ws = wb.create_sheet("Painel do Coordenador")
ws.sheet_view.showGridLines = False
r = sheet_title(
    ws, 1, "Painel do coordenador — leitura mensal de gestão",
    "Referências de ritmo — não são metas de cobrança. Todos os volumes são PARÂMETRO A VALIDAR (faixas de referência do Playbook v1, a calibrar no 1º ciclo). Meta-mãe: ramp 40 → 60 → 80 → 100+ parceiros ativos.",
)
heads = ["Indicador de gestão", "Referência de ritmo (PARÂMETRO A VALIDAR)",
         "Semana 1", "Semana 2", "Semana 3", "Semana 4", "Mês", "Leitura de gestão"]
widths = [34, 30, 10, 10, 10, 10, 12, 40]
header_row(ws, r, heads, widths)
# (indicador, referência, mês é soma? True/False)
inds = [
    ("Parceiros ativos (régua de atividade)", "Ramp trimestral: 40 → 60 → 80 → 100+", False),
    ("Novos parceiros credenciados (Funil 1)", "4–6 / mês, com carteira mapeada", True),
    ("Reativações da base (stand-by → ativo)", "6–8 / mês nos 2 primeiros ciclos", True),
    ("Contatos efetivos de gestão realizados", "20–25 / semana (Tier A quinzenal · B mensal)", True),
    ("Aderência à régua de contato", "≥ 90% dos contatos programados", False),
    ("Ligações de prospecção de novos parceiros", "10–12 / semana", True),
    ("Reuniões de apresentação realizadas", "3–4 / semana", True),
    ("Reuniões de planejamento & resultado", "5–6 / semana", True),
    ("Visitas presenciais", "1–2 / semana", True),
    ("Indicações qualificadas geradas (Funil 2)", "≥ 1 por parceiro ativo / trimestre", True),
    ("Reuniões comerciais geradas ao time de vendas", "Baseline no 1º ciclo (acompanhar)", True),
    ("Tempo credenciamento → 1ª indicação", "≤ 45 dias por parceiro", False),
]
for i, (nome, ref, soma) in enumerate(inds):
    row = r + 1 + i
    ws.row_dimensions[row].height = 30
    vals = [nome, ref, "", "", "", "", None, ""]
    for j, v in enumerate(vals):
        cell = ws.cell(row=row, column=1 + j)
        if j == 6:
            cell.value = f"=IF(COUNT(C{row}:F{row})=0,\"\",SUM(C{row}:F{row}))" if soma else ""
        elif v is not None:
            cell.value = v
        cell.border = BORDER
        cell.alignment = WRAP if j in (0, 1, 7) else WRAP_C
        cell.font = BODY_B if j == 0 else BODY
        if i % 2 == 1:
            cell.fill = ZEBRA
note_row = r + len(inds) + 2
ws.cell(row=note_row, column=1,
        value="Leitura de gestão: o painel responde 'o que a carteira está dizendo?' — ritmo abaixo da referência indica agenda a recompor ou régua a recalibrar, não desempenho a punir. Fontes: Twenty (objeto parcerias), ZapSign, agenda e registro de atividades.").font = SUB_FONT
ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)
ws.cell(row=note_row, column=1).alignment = WRAP
ws.row_dimensions[note_row].height = 30
ws.freeze_panes = f"C{r+1}"

# ------------------------------------------------------------------ ABA 6
ws = wb.create_sheet("Leia-me")
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 110
ws.cell(row=2, column=2, value="Cronograma Semanal de Gestão — Parcerias · Como usar").font = TITLE_FONT
ws.cell(row=3, column=2, value="Ferramenta de planejamento e controle do Coordenador de Parcerias (Ulisses) · Lucro Ativo · v1 — setembro/2026").font = SUB_FONT
bullets = [
    ("Filosofia", "Esta é uma ferramenta de GESTÃO, não de cobrança: ela organiza a semana por temas de gestão (planejar → executar → registrar → analisar) e produz leitura de carteira — não é lista de tarefas nem régua punitiva."),
    ("Semana de Gestão", "É a semana-padrão temática. Use-a como pano de fundo: os blocos indicam o tipo de trabalho de cada faixa de horário; a agenda real de cada semana se monta em cima dela."),
    ("Planner Semanal", "Modelo em branco: duplique a aba a cada semana, preencha o foco do dia, os parceiros do ciclo e as reuniões, e feche a semana com a linha de fechamento (que alimenta a pauta da Weekly com o CCO)."),
    ("Cadência da Carteira", "Controle vivo por parceiro. Preencha Tier e Último contato: o Próximo contato calcula sozinho (A = +14 dias · B = +30 dias · C = régua automatizada). Vermelho = contato vencido · amarelo = vence em até 3 dias. As 5 primeiras linhas são exemplos fictícios — substitua pelos parceiros reais."),
    ("Rituais & Ciclo", "Mapa dos rituais do canal com entrada e saída de cada um — o que alimenta a reunião e o que ela deve produzir. Um ritual sem saída registrada não aconteceu."),
    ("Painel do Coordenador", "Leitura mensal de ritmo. As referências são faixas do Playbook v1 marcadas como PARÂMETRO A VALIDAR — servem para calibrar a agenda no 1º ciclo, não para cobrar meta."),
    ("Registro", "Fonte única de verdade é o Twenty (objeto 'parcerias'): contato efetivo = conversa real com próximo passo datado. Esta planilha organiza e lê; o registro oficial mora no CRM."),
    ("Fonte", "Playbook de Gestão de Parcerias v1.0 (minuta em validação com o Ulisses e o CCO). Revisar esta ferramenta junto com a revisão do playbook — fevereiro/2027 ou antes, ao fim do 1º ciclo."),
]
row = 5
for titulo, txt in bullets:
    c = ws.cell(row=row, column=2, value=f"■  {titulo} — {txt}")
    c.font = BODY
    c.alignment = WRAP
    ws.row_dimensions[row].height = 42
    row += 1
ws.cell(row=row + 1, column=2, value="@meulucroativo · Gestão de Parcerias · Volumes de referência — PARÂMETRO A VALIDAR na revisão com o coordenador.").font = SUB_FONT

OUT = "/home/hermes/agente-rene/work/parcerias_cron/Cronograma_Semanal_Gestao_Parcerias_Ulisses_v1.xlsx"
wb.save(OUT)
print("salvo:", OUT)
