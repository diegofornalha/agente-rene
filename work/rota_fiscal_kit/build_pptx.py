#!/usr/bin/env python3
# Deck Rota Fiscal — PPTX nativo editável (texto real, não imagem).
# Espelha o deck.html: 12 slides, navy #010860 + dourado #C9A227,
# Montserrat (títulos) / Inter (corpo).
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x01, 0x08, 0x60)
NAVY2 = RGBColor(0x03, 0x06, 0x3F)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
INK = RGBColor(0x1A, 0x1F, 0x36)
GREY = RGBColor(0x5A, 0x60, 0x72)
BG = RGBColor(0xF7, 0xF8, 0xFC)
LINE = RGBColor(0xE3, 0xE6, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LILAC = RGBColor(0xC8, 0xCD, 0xF0)
NOTEBG = RGBColor(0xFD, 0xF9, 0xEE)
NOTEBRD = RGBColor(0xEC, 0xD9, 0xA0)
NOTEINK = RGBColor(0x6B, 0x55, 0x10)

SW, SH = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, fill, line=None, line_w=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(0.75)
    sp.shadow.inherit = False
    return sp


def text(s, x, y, w, h, runs, size, color, bold=False, font='Inter',
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, leading=1.15, space_after=0):
    """runs: str ou lista de (texto, bold, cor) por parágrafo; lista de parágrafos."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = leading
        if space_after:
            p.space_after = Pt(space_after)
        if isinstance(para, str):
            para = [(para, bold, color)]
        for txt, b, c in para:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.bold = b
            r.font.color.rgb = c
            r.font.name = font
    return tb


def topbar(s, num, label, page):
    text(s, Inches(0.6), Inches(0.28), Inches(9), Inches(0.4),
         [[(num + '  ', True, GOLD), (label.upper(), True, NAVY)]],
         13, NAVY, font='Montserrat')
    rect(s, Inches(0.6), Inches(0.72), Inches(8.6), Pt(2.4), NAVY)
    rect(s, Inches(9.2), Inches(0.72), Inches(3.5), Pt(2.4), GOLD)
    text(s, Inches(0.6), Inches(7.05), Inches(4), Inches(0.35),
         '@meulucroativo', 11, NAVY, bold=True, font='Montserrat')
    text(s, Inches(12.2), Inches(7.05), Inches(0.6), Inches(0.35),
         page, 11, GREY, align=PP_ALIGN.RIGHT)


def title_sub(s, title, sub, y=1.05):
    text(s, Inches(0.6), Inches(y), Inches(12.1), Inches(0.75), title,
         30, NAVY, bold=True, font='Montserrat')
    text(s, Inches(0.6), Inches(y + 0.72), Inches(12.1), Inches(0.7), sub,
         14.5, GREY, leading=1.25)


def card(s, x, y, w, h, head, body, accent=NAVY, body_size=12.5):
    rect(s, x, y, w, h, BG, LINE)
    rect(s, x, y, Pt(4.5), h, accent)
    text(s, x + Inches(0.18), y + Inches(0.10), w - Inches(0.32), Inches(0.35),
         head, 14, NAVY, bold=True, font='Montserrat')
    text(s, x + Inches(0.18), y + Inches(0.48), w - Inches(0.32), h - Inches(0.6),
         body, body_size, INK, leading=1.22)


def note(s, x, y, w, h, txt):
    rect(s, x, y, w, h, NOTEBG, NOTEBRD)
    text(s, x + Inches(0.15), y + Inches(0.08), w - Inches(0.3), h - Inches(0.16),
         txt, 11.5, NOTEINK, leading=1.2)


def bullet_list(s, x, y, w, items, size=14, gap=0.62):
    for i, (b, rest) in enumerate(items):
        yy = y + Inches(gap * i)
        rect(s, x, yy + Inches(0.08), Inches(0.12), Inches(0.12), GOLD)
        text(s, x + Inches(0.28), yy, w - Inches(0.28), Inches(gap),
             [[(b, True, NAVY), (rest, False, INK)]], size, INK, leading=1.2)


def cover(s, kicker, title_lines, sub, metas, logo_w=3.2):
    rect(s, 0, 0, SW, SH, NAVY)
    rect(s, 0, 0, Inches(0.16), SH, GOLD)
    s.shapes.add_picture('logo_branca.png', Inches(0.95), Inches(0.75),
                         width=Inches(logo_w))
    text(s, Inches(0.95), Inches(2.05), Inches(11), Inches(0.4), kicker.upper(),
         14, GOLD, bold=True, font='Montserrat')
    text(s, Inches(0.95), Inches(2.5), Inches(11.5), Inches(1.8), title_lines,
         40, WHITE, bold=True, font='Montserrat', leading=1.05)
    text(s, Inches(0.95), Inches(4.35), Inches(10.6), Inches(1.3), sub,
         16, LILAC, leading=1.3)
    text(s, Inches(0.95), Inches(6.3), Inches(11.5), Inches(0.5),
         [metas], 13, LILAC)
    # barras decorativas
    for i, hh in enumerate([0.7, 1.1, 1.6, 2.2]):
        rect(s, Inches(10.4 + i * 0.55), SH - Inches(0.6 + hh),
             Inches(0.36), Inches(hh), GOLD)


# ── 1 · CAPA ──
s = slide()
cover(s, 'Solução Internacional · Rota Fiscal',
      ['Internacionalização', 'tributária e patrimonial'],
      'O caminho estruturado para o empresário brasileiro diversificar jurisdição, proteger patrimônio e reduzir a carga tributária — de forma lícita, documentada e com condução jurídica integral da Lucro Ativo.',
      [('Lucro Ativo', True, WHITE), (' · Compliance Fiscal Tributário          Apresentação Comercial', False, LILAC)])

# ── 2 · O CENÁRIO ──
s = slide()
topbar(s, '01', 'O cenário', '02')
title_sub(s, 'O custo de operar apenas no Brasil',
          'O empresário brasileiro carrega uma das maiores cargas tributárias e um dos ambientes fiscais mais complexos do mundo — e a transição da Reforma Tributária (EC 132/2023) adiciona anos de convivência entre dois sistemas.')
cw, chh = Inches(6.0), Inches(1.95)
card(s, Inches(0.6), Inches(2.75), cw, chh, 'Concentração de risco',
     'Patrimônio, faturamento e residência fiscal em uma única jurisdição significam exposição total a mudanças de regra, contencioso e instabilidade — sem alternativa preparada.')
card(s, Inches(6.75), Inches(2.75), cw, chh, 'Carga sobre o resultado',
     'Entre tributos sobre consumo, renda e folha, a parcela do resultado entregue ao fisco brasileiro compromete a capacidade de reinvestimento e a competitividade internacional.')
card(s, Inches(0.6), Inches(4.85), cw, chh, 'Complexidade crescente',
     'A transição da Reforma (2026–2033) impõe dois sistemas em paralelo — CBS/IBS convivendo com o modelo atual — elevando o custo de conformidade justamente para quem produz.')
card(s, Inches(6.75), Inches(4.85), cw, chh, 'A resposta racional',
     'Internacionalizar não é fugir: é planejamento tributário lícito — escolher, dentro da lei, a jurisdição e a estrutura mais eficientes para cada parte do negócio e do patrimônio.',
     accent=GOLD)

# ── 3 · A SOLUÇÃO ──
s = slide()
topbar(s, '02', 'A solução', '03')
title_sub(s, 'Rota Fiscal: internacionalização com método',
          'Um programa estruturado que leva o empresário do diagnóstico à operação internacional ativa — com advocacia tributária na condução e estrutura própria da Lucro Ativo no Paraguai.')
bullet_list(s, Inches(0.6), Inches(2.8), Inches(12.1), [
    ('Sob medida: ', 'a rota é definida pelo perfil — onde está a receita, o patrimônio, a família e o plano de vida do empresário. Não existe pacote único.'),
    ('Condução integral: ', 'diagnóstico, desenho da estrutura, documentação, atos no exterior e manutenção — um único responsável do início ao fim.'),
    ('Compliance nos dois lados: ', 'cada passo é formalizado tanto perante o fisco brasileiro quanto perante a jurisdição de destino. Nada de atalhos.'),
    ('Presença local: ', 'operação própria em Assunção (EAS Paraguay), com equipe e parceiros credenciados no destino — não terceirizamos o que é crítico.'),
], gap=0.7)
note(s, Inches(0.6), Inches(5.95), Inches(12.1), Inches(0.72),
     [[('Posicionamento: ', True, NOTEINK), ('a Rota Fiscal é planejamento tributário internacional lícito. A Lucro Ativo não realiza — nem orienta — ocultação de patrimônio ou omissão de declarações.', False, NOTEINK)]])

# ── 4 · AS TRÊS ROTAS ──
s = slide()
topbar(s, '03', 'As três rotas', '04')
title_sub(s, 'Três caminhos, um método',
          'O diagnóstico define qual rota — ou combinação de rotas — serve ao seu caso.')
rotas = [
    ('ROTA 01', 'Paraguai',
     'Residência fiscal e estrutura empresarial no país vizinho, com tributação radicalmente mais simples.',
     ['Residência e RUC no Paraguai', 'Regime maquila para produzir e exportar', 'Mercosul: logística e mercado'],
     [('Para quem busca ', False, GREY), ('mudança de base fiscal', True, NAVY), (' ou produção no Mercosul.', False, GREY)]),
    ('ROTA 02', 'LLC · EUA',
     'Veículo empresarial norte-americano para receitas internacionais e organização patrimonial em moeda forte.',
     ['Faturamento internacional em dólar', 'Acesso ao sistema financeiro dos EUA', 'Declarada e conforme no Brasil'],
     [('Para quem tem ', False, GREY), ('receita internacional', True, NAVY), (' ou busca dolarização estruturada.', False, GREY)]),
    ('ROTA 03', 'REINTEGRA',
     'Recuperação do resíduo tributário embutido na cadeia de produtos exportados — sem sair do Brasil.',
     ['Crédito sobre a receita de exportação', 'Procedimento administrativo previsto em lei', 'Complementa as demais rotas'],
     [('Para ', False, GREY), ('indústrias exportadoras', True, NAVY), (' que deixam dinheiro na mesa.', False, GREY)]),
]
rx = [Inches(0.6), Inches(4.85), Inches(9.1)]
rw = Inches(4.05)
for (k, t, desc, its, foot), x in zip(rotas, rx):
    rect(s, x, Inches(2.55), rw, Inches(4.35), WHITE, LINE)
    rect(s, x, Inches(2.55), rw, Inches(0.95), NAVY)
    text(s, x + Inches(0.18), Inches(2.63), rw - Inches(0.3), Inches(0.3), k, 10.5, GOLD, bold=True, font='Montserrat')
    text(s, x + Inches(0.18), Inches(2.92), rw - Inches(0.3), Inches(0.45), t, 19, WHITE, bold=True, font='Montserrat')
    text(s, x + Inches(0.18), Inches(3.62), rw - Inches(0.34), Inches(1.0), desc, 12, INK, leading=1.25)
    for i, it in enumerate(its):
        yy = Inches(4.72 + i * 0.44)
        rect(s, x + Inches(0.18), yy + Inches(0.06), Inches(0.1), Inches(0.1), GOLD)
        text(s, x + Inches(0.4), yy, rw - Inches(0.55), Inches(0.44), it, 11.5, INK, leading=1.1)
    rect(s, x, Inches(6.15), rw, Inches(0.75), BG, LINE)
    text(s, x + Inches(0.18), Inches(6.24), rw - Inches(0.34), Inches(0.6), [foot], 11, GREY, leading=1.15)

# ── 5 · ROTA PARAGUAI ──
s = slide()
topbar(s, '04', 'Rota Paraguai', '05')
title_sub(s, 'Paraguai: simplicidade tributária ao lado de casa',
          'Um dos sistemas tributários mais enxutos das Américas, a poucas horas do Brasil e dentro do Mercosul.')
nums = [('10%', 'Imposto de renda empresarial (IRE)', NAVY),
        ('10%', 'IVA — imposto sobre valor agregado', NAVY),
        ('1%', 'Tributo único do regime maquila sobre o valor agregado local', GOLD)]
for i, (v, l, c) in enumerate(nums):
    x = Inches(0.6 + i * 4.25)
    rect(s, x, Inches(2.6), Inches(4.0), Inches(1.35), BG, LINE)
    text(s, x, Inches(2.72), Inches(4.0), Inches(0.6), v, 30, c, bold=True,
         font='Montserrat', align=PP_ALIGN.CENTER)
    text(s, x + Inches(0.2), Inches(3.34), Inches(3.6), Inches(0.55), l, 11, GREY,
         align=PP_ALIGN.CENTER, leading=1.1)
card(s, Inches(0.6), Inches(4.2), Inches(6.0), Inches(1.75), 'Residência fiscal',
     'Processo formal de residência, cédula e registro tributário (RUC), com planejamento da saída fiscal brasileira quando aplicável — sem improviso migratório.')
card(s, Inches(6.75), Inches(4.2), Inches(6.0), Inches(1.75), 'Regime maquila',
     'Produzir no Paraguai e exportar — inclusive para o Brasil — com tributo único de 1% sobre o valor agregado local (Lei 1.064/97), energia barata e mão de obra competitiva.')
note(s, Inches(0.6), Inches(6.15), Inches(12.1), Inches(0.72),
     [[('PARÂMETRO A VALIDAR: ', True, NOTEINK), ('alíquotas, requisitos migratórios e enquadramento no regime maquila são confirmados caso a caso pelo jurídico antes de qualquer proposta.', False, NOTEINK)]])

# ── 6 · ROTA LLC ──
s = slide()
topbar(s, '05', 'Rota LLC · EUA', '06')
title_sub(s, 'LLC: sua operação internacional em dólar',
          'Um veículo empresarial nos Estados Unidos para faturar clientes no exterior, acessar o sistema financeiro americano e organizar patrimônio em moeda forte — declarado e em conformidade no Brasil.')
card(s, Inches(0.6), Inches(2.85), cw, Inches(1.9), 'Para que serve',
     'Receitas de exportação de serviços, marketplaces e clientes internacionais entram direto na estrutura americana, com conta bancária nos EUA e recebimento em dólar.')
card(s, Inches(6.75), Inches(2.85), cw, Inches(1.9), 'Proteção e organização',
     'Separação entre o patrimônio pessoal e a operação, sucessão facilitada e diversificação cambial — o caixa deixa de depender de uma única moeda.')
card(s, Inches(0.6), Inches(4.95), cw, Inches(1.9), 'Conformidade brasileira',
     'A estrutura é declarada no Brasil (DIRPF/CBE) e desenhada à luz das regras de tributação de investimentos no exterior (Lei 14.754/2023). Eficiência sim; omissão jamais.',
     accent=GOLD)
card(s, Inches(6.75), Inches(4.95), cw, Inches(1.9), 'Quando combinar rotas',
     'A LLC convive com a Rota Paraguai: cada peça cuida de uma função — receita internacional, base fiscal e produção — no mesmo desenho.')

# ── 7 · REINTEGRA ──
s = slide()
topbar(s, '06', 'Rota REINTEGRA', '07')
title_sub(s, 'REINTEGRA: o crédito que a exportadora esquece',
          'O regime devolve ao exportador parte do resíduo tributário acumulado na cadeia produtiva — um direito previsto em lei (Lei 13.043/2014) que muitas indústrias simplesmente não exercem.')
kvs = [('1', 'Quem alcança: ', 'indústrias que exportam bens manufaturados elegíveis — diretamente ou via trading.'),
       ('2', 'Como funciona: ', 'apuração de crédito sobre a receita de exportação, aproveitável para compensar tributos federais ou pedir ressarcimento — procedimento administrativo, sem litígio.'),
       ('3', 'Efeito caixa: ', 'recupera margem sem mudar a operação — e financia, muitas vezes, o próprio projeto de internacionalização.')]
for i, (n, b, rest) in enumerate(kvs):
    y = Inches(2.95 + i * 0.95)
    rect(s, Inches(0.6), y, Inches(0.5), Inches(0.5), NAVY)
    text(s, Inches(0.6), y + Inches(0.04), Inches(0.5), Inches(0.42), n, 16, WHITE,
         bold=True, font='Montserrat', align=PP_ALIGN.CENTER)
    text(s, Inches(1.3), y, Inches(11.4), Inches(0.9),
         [[(b, True, NAVY), (rest, False, INK)]], 14, INK, leading=1.25)
note(s, Inches(0.6), Inches(6.0), Inches(12.1), Inches(0.85),
     [[('PARÂMETRO A VALIDAR: ', True, NOTEINK), ('o percentual do REINTEGRA é fixado por decreto e varia no tempo; a alíquota aplicável e o alcance retroativo são confirmados pelo jurídico antes de qualquer projeção ao cliente.', False, NOTEINK)]])

# ── 8 · JORNADA ──
s = slide()
topbar(s, '07', 'Como funciona', '08')
title_sub(s, 'A jornada em seis etapas',
          'Método testado na prática — cada etapa com responsável, prazo e evidência documental.')
steps = [
    ('1', 'Diagnóstico', 'Raio-X tributário e patrimonial: situação fiscal na Receita, DIRPF, estrutura societária e objetivos do empresário.', False),
    ('2', 'Desenho da rota', 'Definição da rota (ou combinação) e da modalidade de residência/estrutura, com parecer e cronograma.', False),
    ('3', 'Documentação', 'Certidões, autenticações, apostilamento de Haia e tradução juramentada — o dossiê completo antes de qualquer ato.', False),
    ('4', 'Planejamento da saída', 'Quando aplicável: regularização de DIRPF e Declaração de Saída Definitiva (DSDP), para não haver bitributação nem pendência.', False),
    ('5', 'Missão presencial', 'Agenda organizada no destino — migrações, registros e abertura da estrutura — com acompanhamento local da equipe.', True),
    ('6', 'Ativação e manutenção', 'RUC/registros ativos, contabilidade local e calendário de obrigações — a estrutura viva, não um carimbo no passaporte.', False),
]
for i, (n, t, d, hot) in enumerate(steps):
    x = Inches(0.6 + (i % 3) * 4.25)
    y = Inches(2.65 + (i // 3) * 2.15)
    rect(s, x, y, Inches(4.0), Inches(1.95), NOTEBG if hot else BG, LINE)
    rect(s, x, y, Inches(4.0), Pt(4.5), GOLD if hot else NAVY)
    text(s, x + Inches(0.15), y + Inches(0.08), Inches(1), Inches(0.45), n, 22,
         GOLD, bold=True, font='Montserrat')
    text(s, x + Inches(0.15), y + Inches(0.52), Inches(3.7), Inches(0.35), t,
         13.5, NAVY, bold=True, font='Montserrat')
    text(s, x + Inches(0.15), y + Inches(0.88), Inches(3.72), Inches(1.0), d,
         10.5, GREY, leading=1.2)

# ── 9 · SEGURANÇA ──
s = slide()
topbar(s, '08', 'Segurança jurídica', '09')
title_sub(s, 'Tudo dentro da lei — nos dois países',
          'A diferença entre planejamento e problema é a formalização. A Rota Fiscal existe para fazer do jeito certo.')
card(s, Inches(0.6), Inches(2.85), cw, Inches(1.9), 'Saída fiscal formalizada',
     'Quem muda de base fiscal faz a Declaração de Saída Definitiva e encerra as obrigações corretamente — sem virar residente "fantasma" em dois países.')
card(s, Inches(6.75), Inches(2.85), cw, Inches(1.9), 'Estruturas declaradas',
     'Participações no exterior são declaradas no Brasil (DIRPF/CBE) enquanto houver residência fiscal brasileira. Nenhuma estrutura nossa depende de omissão.')
card(s, Inches(0.6), Inches(4.95), cw, Inches(1.9), 'Advocacia na condução',
     'Cada caso tem análise jurídica individual — advocacia tributária com 15 anos de prática, não um curso ou um despachante.')
card(s, Inches(6.75), Inches(4.95), cw, Inches(1.9), 'O que não fazemos',
     'Não prometemos "imposto zero", não operamos ocultação de ativos e não executamos rota sem diagnóstico. Se não servir ao seu caso, dizemos.',
     accent=GOLD)

# ── 10 · POR QUE A LUCRO ATIVO ──
s = slide()
topbar(s, '09', 'Por que a Lucro Ativo', '10')
title_sub(s, 'Estrutura própria, não intermediação',
          'A maioria vende consultoria à distância. A Lucro Ativo opera no destino.')
bullet_list(s, Inches(0.6), Inches(2.85), Inches(12.1), [
    ('Presença em Assunção: ', 'estrutura própria no Paraguai (EAS Paraguay), com equipe local e rede credenciada — despachos, registros e contabilidade sem depender de terceiros desconhecidos.'),
    ('Base tributária brasileira: ', 'quem desenha a rota entende profundamente o lado de cá — Receita Federal, DIRPF, saída fiscal e as regras de tributação internacional do Brasil.'),
    ('Ecossistema completo: ', 'do diagnóstico à manutenção anual, o mesmo time responde — com CRM, documentação e trilha de auditoria de cada etapa.'),
    ('Portfólio integrado: ', 'a Rota Fiscal conversa com as demais soluções da casa (créditos tributários, compensação de ICMS, transação federal) — o empresário resolve o Brasil enquanto constrói o internacional.'),
], gap=0.95)

# ── 11 · PARA QUEM É ──
s = slide()
topbar(s, '10', 'Para quem é', '11')
title_sub(s, 'O perfil que mais aproveita a Rota Fiscal',
          'Se você marca dois ou mais pontos abaixo, o diagnóstico tende a se pagar.')
perfis = [
    ('Empresário de resultado', 'Lucro consistente e sensação de que o fisco é o sócio majoritário — sem plano B de jurisdição.', NAVY),
    ('Exportador / receita externa', 'Indústria que exporta ou prestador com clientes internacionais recebendo em estrutura brasileira.', NAVY),
    ('Patrimônio a proteger', 'Ativos relevantes concentrados no Brasil e sucessão ainda não resolvida.', NAVY),
    ('Produção intensiva', 'Operação industrial pressionada por custo tributário, energético e de folha — candidata natural à maquila.', NAVY),
    ('Mobilidade real', 'Disposição de cumprir os requisitos formais de residência e presença que a rota escolhida exigir.', NAVY),
    ('Visão de longo prazo', 'Entende que internacionalizar é projeto estruturado — não atalho de fim de ano fiscal.', GOLD),
]
for i, (t, d, acc) in enumerate(perfis):
    x = Inches(0.6 + (i % 3) * 4.25)
    y = Inches(2.75 + (i // 3) * 2.05)
    card(s, x, y, Inches(4.0), Inches(1.85), t, d, accent=acc, body_size=11.5)

# ── 12 · PRÓXIMOS PASSOS ──
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, 0, Inches(0.16), SH, GOLD)
s.shapes.add_picture('logo_branca.png', Inches(0.95), Inches(0.85), width=Inches(2.6))
text(s, Inches(0.95), Inches(2.0), Inches(11), Inches(0.4), 'PRÓXIMOS PASSOS',
     14, GOLD, bold=True, font='Montserrat')
text(s, Inches(0.95), Inches(2.45), Inches(11.5), Inches(0.9),
     'Comece pelo diagnóstico', 36, WHITE, bold=True, font='Montserrat')
text(s, Inches(0.95), Inches(3.45), Inches(10.4), Inches(1.4),
     'Uma reunião de diagnóstico define se existe rota para o seu caso, qual é ela e o que muda no seu resultado. Sem compromisso de execução — e com a resposta honesta, inclusive quando a resposta for "ainda não".',
     16, LILAC, leading=1.3)
text(s, Inches(0.95), Inches(5.15), Inches(11.5), Inches(0.5),
     [[('1. ', True, GOLD), ('Reunião de diagnóstico      ', False, WHITE),
       ('2. ', True, GOLD), ('Parecer e desenho da rota      ', False, WHITE),
       ('3. ', True, GOLD), ('Execução acompanhada', False, WHITE)]], 15, WHITE)
text(s, Inches(0.95), Inches(6.35), Inches(11.5), Inches(0.5),
     [[('Lucro Ativo', True, WHITE), (' · Compliance Fiscal Tributário · @meulucroativo', False, LILAC)]],
     13, LILAC)
for i, hh in enumerate([0.7, 1.1, 1.6, 2.2]):
    rect(s, Inches(10.4 + i * 0.55), SH - Inches(0.6 + hh), Inches(0.36), Inches(hh), GOLD)

prs.save('Deck_Comercial_RotaFiscal_LucroAtivo_v1.pptx')
print('PPTX salvo:', prs.slide_width, 'x', prs.slide_height, '-', len(prs.slides.__iter__.__self__._sldIdLst), 'slides')
