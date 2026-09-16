#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""FAQ de Capacitação Comercial — Rota Fiscal (Lucro Ativo), padrão do FAQ v1.3."""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x01, 0x08, 0x60)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
INK = RGBColor(0x14, 0x14, 0x14)
GREY = RGBColor(0x5A, 0x60, 0x72)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD_DARK = RGBColor(0x8A, 0x6D, 0x0D)

doc = Document()

# ---------- estilos base ----------
st = doc.styles['Normal']
st.font.name = 'Inter'; st.font.size = Pt(10.5); st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.25
for name, sz, before in (('Heading 1', 20, 18), ('Heading 2', 15, 12), ('Heading 3', 12.5, 8)):
    h = doc.styles[name]
    h.font.name = 'Montserrat'; h.font.size = Pt(sz); h.font.bold = True
    h.font.color.rgb = NAVY
    h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(6)
    rf = h.element.rPr.get_or_add_rFonts()
    for attr in ('asciiTheme', 'hAnsiTheme', 'cstheme', 'eastAsiaTheme'):
        if rf.get(qn('w:' + attr)) is not None:
            del rf.attrib[qn('w:' + attr)]
    rf.set(qn('w:ascii'), 'Montserrat'); rf.set(qn('w:hAnsi'), 'Montserrat'); rf.set(qn('w:cs'), 'Montserrat')
lb = doc.styles['List Bullet']
lb.font.name = 'Inter'; lb.font.size = Pt(10.5)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hexcolor); tcPr.append(shd)

def no_borders(table):
    tbl = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'none'); borders.append(e)
    tbl.append(borders)

def set_borders(table, color='E3E6EF', sz=4):
    tbl = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz)); e.set(qn('w:color'), color); borders.append(e)
    tbl.append(borders)

def add_runs(p, text, bold_color=NAVY, base_color=None, size=None, font='Inter'):
    """**bold** vira negrito navy."""
    for i, part in enumerate(re.split(r'\*\*', text)):
        if not part:
            continue
        r = p.add_run(part)
        r.font.name = font
        if size: r.font.size = Pt(size)
        if i % 2 == 1:
            r.font.bold = True; r.font.color.rgb = bold_color
        elif base_color is not None:
            r.font.color.rgb = base_color
    return p

# ---------- CAPA (full-bleed navy) ----------
sec = doc.sections[0]
sec.page_width = Cm(21); sec.page_height = Cm(29.7)
for a in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
    setattr(sec, a, Cm(0))

cover = doc.add_table(rows=1, cols=1)
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
cover.autofit = False
cover.columns[0].width = Cm(21)
cover.rows[0].cells[0].width = Cm(21)
tr = cover.rows[0]._tr
trPr = tr.get_or_add_trPr()
trH = OxmlElement('w:trHeight'); trH.set(qn('w:val'), str(int(29.6 * 567))); trH.set(qn('w:hRule'), 'exact')
trPr.append(trH)
no_borders(cover)
c = cover.rows[0].cells[0]
shade(c, '010860')
c.paragraphs[0].text = ''
# espaçador superior
for _ in range(6):
    sp = c.add_paragraph(); sp.paragraph_format.space_after = Pt(18)
plogo = c.add_paragraph(); plogo.paragraph_format.left_indent = Cm(2.4)
plogo.add_run().add_picture('logo_branca.png', width=Cm(7.2))
k = c.add_paragraph(); k.paragraph_format.left_indent = Cm(2.4); k.paragraph_format.space_before = Pt(28)
r = k.add_run('FAQ DE CAPACITAÇÃO COMERCIAL'); r.font.name = 'Montserrat'; r.font.bold = True
r.font.size = Pt(12); r.font.color.rgb = GOLD
t = c.add_paragraph(); t.paragraph_format.left_indent = Cm(2.4); t.paragraph_format.space_before = Pt(6)
r = t.add_run('Rota Fiscal'); r.font.name = 'Montserrat'; r.font.bold = True; r.font.size = Pt(40); r.font.color.rgb = WHITE
t2 = c.add_paragraph(); t2.paragraph_format.left_indent = Cm(2.4)
r = t2.add_run('Internacionalização tributária e patrimonial'); r.font.name = 'Montserrat'
r.font.size = Pt(17); r.font.color.rgb = RGBColor(0xC8, 0xCD, 0xF0)
sub = c.add_paragraph(); sub.paragraph_format.left_indent = Cm(2.4); sub.paragraph_format.space_before = Pt(22)
r = sub.add_run('Guia do executivo de vendas — Paraguai · LLC EUA · REINTEGRA')
r.font.name = 'Inter'; r.font.size = Pt(12); r.font.color.rgb = RGBColor(0xAA, 0xB1, 0xE8)
for _ in range(5):
    sp = c.add_paragraph(); sp.paragraph_format.space_after = Pt(18)
ft = c.add_paragraph(); ft.paragraph_format.left_indent = Cm(2.4)
r = ft.add_run('Lucro Ativo · Compliance Fiscal Tributário · @meulucroativo')
r.font.name = 'Montserrat'; r.font.size = Pt(10.5); r.font.color.rgb = WHITE
v = c.add_paragraph(); v.paragraph_format.left_indent = Cm(2.4)
r = v.add_run('Material interno — não circular para clientes · Versão 1.0 · Setembro de 2026')
r.font.name = 'Inter'; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0xAA, 0xB1, 0xE8)

# ---------- SEÇÃO DE CONTEÚDO ----------
sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
sec2.page_width = Cm(21); sec2.page_height = Cm(29.7)
sec2.top_margin = Cm(2.2); sec2.bottom_margin = Cm(2.0)
sec2.left_margin = Cm(2.4); sec2.right_margin = Cm(2.4)
sec2.footer.is_linked_to_previous = False
sec2.header.is_linked_to_previous = False
fp = sec2.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run('@meulucroativo  ·  FAQ Rota Fiscal — uso interno  ·  pág. ')
r.font.name = 'Inter'; r.font.size = Pt(8.5); r.font.color.rgb = GREY
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
rr = OxmlElement('w:r'); rp = OxmlElement('w:rPr')
sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '17'); rp.append(sz)
col = OxmlElement('w:color'); col.set(qn('w:val'), '5A6072'); rp.append(col)
rr.append(rp); tnode = OxmlElement('w:t'); tnode.text = '1'; rr.append(tnode)
fld.append(rr); fp._p.append(fld)

def H1(txt): doc.add_heading(txt, level=1)
def H2(txt): doc.add_heading(txt, level=2)

def para(txt, size=10.5, space_after=6):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    add_runs(p, txt, size=size)
    return p

def bullet(txt):
    p = doc.add_paragraph(style='List Bullet')
    add_runs(p, txt)
    return p

def question(n, q):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run('▸  %d. ' % n); r.font.name = 'Montserrat'; r.font.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = GOLD_DARK
    r = p.add_run(q); r.font.name = 'Montserrat'; r.font.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = NAVY

def answer(txt):
    para(txt)

def callout(title, body, fill='FDF9EE', border='C9A227', title_color=GOLD_DARK):
    tb = doc.add_table(rows=1, cols=1)
    tb.autofit = False
    tb.columns[0].width = Cm(16.2); tb.rows[0].cells[0].width = Cm(16.2)
    set_borders(tb, color=border, sz=6)
    cell = tb.rows[0].cells[0]; shade(cell, fill)
    p0 = cell.paragraphs[0]
    r = p0.add_run(title); r.font.name = 'Montserrat'; r.font.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = title_color
    p1 = cell.add_paragraph(); add_runs(p1, body, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ---------- FICHA ----------
H2('Ficha do Documento')
ficha = [
    ('Título', 'FAQ de Capacitação Comercial — Rota Fiscal (Internacionalização Tributária e Patrimonial)'),
    ('Natureza', 'Material de capacitação interna (não é material de cliente)'),
    ('Produto coberto', 'Rota Fiscal — Rota Paraguai (residência fiscal e maquila) · Rota LLC EUA · Rota REINTEGRA'),
    ('Finalidade', 'Treinar o time comercial — base técnica e jurídica para conduzir a venda e responder objeções'),
    ('Próximo passo', 'Após validação técnica/jurídica, gerar as versões com foco no cliente'),
    ('Versão / Data', 'Versão 1.0 · Setembro de 2026'),
    ('Base normativa', 'Legislação citada ao fim do documento — todo parâmetro comercial deve ser validado pelo jurídico'),
]
tb = doc.add_table(rows=len(ficha) + 1, cols=2)
tb.autofit = False
set_borders(tb)
tb.columns[0].width = Cm(3.8); tb.columns[1].width = Cm(12.4)
hdr = tb.rows[0]
for j, txt in enumerate(('Campo', 'Conteúdo')):
    cell = hdr.cells[j]; cell.width = (Cm(3.8), Cm(12.4))[j]; shade(cell, '010860')
    r = cell.paragraphs[0].add_run(txt); r.font.name = 'Montserrat'; r.font.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = WHITE
for i, (kk, vv) in enumerate(ficha):
    row = tb.rows[i + 1]
    row.cells[0].width = Cm(3.8); row.cells[1].width = Cm(12.4)
    if i % 2 == 1:
        shade(row.cells[0], 'F2F4FB'); shade(row.cells[1], 'F2F4FB')
    r = row.cells[0].paragraphs[0].add_run(kk); r.font.name = 'Inter'; r.font.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = NAVY
    r = row.cells[1].paragraphs[0].add_run(vv); r.font.name = 'Inter'; r.font.size = Pt(9.5)
doc.add_paragraph().paragraph_format.space_after = Pt(2)

callout('▸ REGRA DE OURO DA CAPACITAÇÃO',
        'Todo número comercial — alíquota, percentual de economia, prazo, honorário, ROI — é **PARÂMETRO A VALIDAR** com o jurídico/tributário antes de chegar ao cliente. O que se garante é **rigor técnico e condução especializada** — nunca resultado, prazo ou percentual. Em internacionalização isso vale em dobro: regras migratórias e tributárias de outro país mudam, e quem promete número vira réu da própria promessa.')

para('**COMO ESTE GUIA ESTÁ ORGANIZADO** — Bloco 1: conceito e posicionamento. Blocos 2 a 4: uma rota por bloco (Paraguai, LLC EUA, REINTEGRA). Bloco 5: jornada e operação. Bloco 6: objeções frequentes com a resposta pronta. Ao final, a base normativa de referência.', space_after=10)

# ---------- SUMÁRIO ----------
H1('Sumário')
p = doc.add_paragraph()
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), r'TOC \o "1-1" \h \z \u')
rr = OxmlElement('w:r'); tnode = OxmlElement('w:t')
tnode.text = 'Sumário automático — no Word, clique com o botão direito e selecione "Atualizar campo".'
rr.append(tnode); fld.append(rr); p._p.append(fld)

# ---------- BLOCO 1 ----------
H1('BLOCO 1 — CONCEITO E POSICIONAMENTO')
question(1, 'O que é, em uma frase, a Rota Fiscal da Lucro Ativo?')
answer('É o **programa de internacionalização tributária e patrimonial** da casa: um método que leva o empresário do **diagnóstico** à **operação internacional ativa** — pela Rota Paraguai, pela Rota LLC EUA, pela Rota REINTEGRA ou por uma combinação delas — com advocacia tributária na condução e estrutura própria da Lucro Ativo em Assunção (EAS Paraguay).')
question(2, 'Rota Fiscal é "ir embora do Brasil"?')
answer('Não necessariamente. Só a mudança de **base fiscal pessoal** (Rota Paraguai com residência) envolve sair formalmente do sistema brasileiro. A **LLC** convive com a residência fiscal no Brasil — declarada e conforme — e o **REINTEGRA** acontece inteiro dentro do Brasil. O diagnóstico define o desenho; não existe pacote único.')
question(3, 'Isso é legal? Como o executivo responde essa pergunta?')
answer('Com tranquilidade: a Rota Fiscal é **planejamento tributário lícito**. Cada passo é formalizado perante o fisco brasileiro e perante a jurisdição de destino — saída fiscal declarada, estruturas declaradas (DIRPF/CBE), contabilidade local ativa. O posicionamento é textual: a Lucro Ativo **não realiza nem orienta** ocultação de patrimônio ou omissão de declarações.')
question(4, 'Qual a diferença entre o que fazemos e evasão fiscal?')
answer('**Elisão** é escolher, dentro da lei, a jurisdição e a estrutura mais eficientes — direito de qualquer empresário. **Evasão** é ocultar patrimônio, omitir declaração, manter residência "fantasma" em dois países. A Rota Fiscal existe exatamente para o empresário internacionalizar **do jeito certo** — com documento, prazo e evidência de cada ato.')
question(5, 'O que a Lucro Ativo NÃO promete?')
answer('Não prometemos **"imposto zero"**, não operamos ocultação de ativos e não executamos rota sem diagnóstico. Se a internacionalização não servir ao caso do cliente, a resposta é honesta: **"ainda não"**. Essa postura é argumento de venda — diferencia a casa de curso de guru e de despachante.')

# ---------- BLOCO 2 ----------
H1('BLOCO 2 — ROTA PARAGUAI')
question(6, 'Por que o Paraguai?')
answer('Um dos sistemas tributários mais enxutos das Américas, ao lado de casa e dentro do Mercosul: **IRE de 10%** (imposto de renda empresarial), **IVA de 10%** e, para a indústria, o **regime maquila** com tributo único de **1%** sobre o valor agregado local ou sobre a fatura de exportação (Lei 7.547/2025, que substituiu a Lei 1.064/97). Soma-se energia barata, mão de obra competitiva e logística de vizinho. **PARÂMETRO A VALIDAR:** alíquotas e enquadramentos são confirmados caso a caso antes de qualquer proposta.')
question(7, 'O que é o regime maquila?')
answer('Regime da **Lei paraguaia 7.547/2025** (regulamentada pelo Decreto 5.714/2026; substitui a Lei 1.064/97) que permite instalar produção — e, agora expressamente, **serviços** de exportação como software, suporte e BPO — no Paraguai para exportar, inclusive de volta ao Brasil, pagando **tributo único de 1%** sobre o valor agregado no país ou sobre a fatura de exportação, o que for maior. O programa é aprovado pelo CNIME e a operação precisa ter substância local (pessoal, instalações, cadeia de valor). É a rota natural da indústria e do serviço pressionados por custo tributário, energético e de folha. O enquadramento do projeto no regime é análise jurídica, não checkbox.')
question(8, 'O cliente precisa morar no Paraguai?')
answer('Depende do objetivo. **Mudança de base fiscal pessoal** exige cumprir os requisitos formais de residência e formalizar a saída do Brasil. **Estrutura empresarial** (maquila, holding local) é outro desenho, que pode existir sem mudança pessoal. O que não existe é atalho: residência de papel, sem substância, é o erro que a Rota Fiscal foi desenhada para evitar.')
question(9, 'O que é RUC?')
answer('**Registro Único do Contribuinte** — o registro tributário paraguaio, análogo à inscrição CPF/CNPJ na Receita. Ter RUC ativo, contabilidade local e obrigações em dia é o que transforma a estrutura em **operação viva**, não um carimbo no passaporte.')
question(10, 'O que acontece com as obrigações do cliente no Brasil?')
answer('Quem muda de base fiscal faz a **Declaração de Saída Definitiva** e encerra as obrigações corretamente — sem virar residente "fantasma". Quem mantém residência brasileira **declara** as participações no exterior (DIRPF/CBE) enquanto ela durar. Nos dois cenários, a regra é a mesma: **nada depende de omissão**.')
question(11, 'Qual o papel da EAS Paraguay?')
answer('É a **estrutura própria da Lucro Ativo em Assunção** — equipe local e rede credenciada para despachos, registros, contabilidade e a missão presencial do cliente. É o diferencial contra a concorrência que vende consultoria à distância: **a Lucro Ativo opera no destino**.')

# ---------- BLOCO 3 ----------
H1('BLOCO 3 — ROTA LLC · EUA')
question(12, 'O que é a LLC e para que ela serve?')
answer('Um **veículo empresarial norte-americano** (Limited Liability Company) para faturar clientes no exterior, acessar o sistema financeiro dos EUA e organizar patrimônio em **moeda forte**. Receitas de exportação de serviços, marketplaces e clientes internacionais entram direto na estrutura americana, com conta nos EUA e recebimento em dólar.')
question(13, 'Para quem a LLC faz sentido?')
answer('Prestadores e empresas com **receita internacional** recebendo em estrutura brasileira; empresários buscando **dolarização estruturada** do caixa; e famílias organizando **sucessão e separação patrimonial** — o patrimônio pessoal deixa de se confundir com a operação e de depender de uma única moeda.')
question(14, 'LLC é "offshore escondida"?')
answer('Não — e essa resposta deve ser dada com ênfase. A estrutura é **declarada no Brasil** (DIRPF/CBE) e desenhada à luz das regras brasileiras de tributação de investimentos no exterior (**Lei 14.754/2023**). Eficiência sim; omissão jamais. Estrutura que depende de o fisco não ver não é planejamento, é passivo.')
question(15, 'Como a LLC é tributada no Brasil?')
answer('Pelas regras da **Lei 14.754/2023** para entidades controladas no exterior — que definem quando o lucro é tributado e as opções de regime disponíveis. Regra geral: o lucro da controlada é tributado no IRPF do sócio a **15%** em 31/12 de cada ano, distribuído ou não; e a **SC COSIT 56/2026** enquadra a LLC transparente de sócio não residente nos EUA como regime fiscal privilegiado — mais um motivo para o desenho ser declarado desde o início. O enquadramento ideal depende do perfil (tipo de receita, distribuição, residência) e **é definido no parecer jurídico do caso**. O executivo não simula tributação de LLC em reunião: **PARÂMETRO A VALIDAR**, sempre.')
question(16, 'A LLC combina com a Rota Paraguai?')
answer('Sim — e essa é a sofisticação do método: cada peça cuida de uma função no mesmo desenho. **Receita internacional** na LLC, **base fiscal e produção** no Paraguai, cada uma declarada onde deve. O diagnóstico define se o caso pede uma rota ou a combinação.')

# ---------- BLOCO 4 ----------
H1('BLOCO 4 — ROTA REINTEGRA')
question(17, 'O que é o REINTEGRA?')
answer('Regime da **Lei 13.043/2014** que devolve ao exportador parte do **resíduo tributário** acumulado na cadeia produtiva dos bens exportados — tributos que se incorporaram ao custo e não foram recuperados. É um direito previsto em lei que muitas indústrias simplesmente **não exercem**.')
question(18, 'Quem tem direito?')
answer('**Indústrias que exportam bens manufaturados elegíveis** — diretamente ou por meio de trading. A elegibilidade depende do produto e da operação; a triagem é parte do diagnóstico.')
question(19, 'Como o crédito é aproveitado?')
answer('Apuração de crédito sobre a **receita de exportação**, aproveitável para **compensar tributos federais** ou para **pedido de ressarcimento** — procedimento administrativo, sem litígio. O efeito é caixa: recupera margem sem mudar a operação, e muitas vezes financia o próprio projeto de internacionalização.')
question(20, 'Qual o percentual do REINTEGRA?')
answer('O percentual é **fixado por decreto e varia no tempo**. Hoje: **0,1%** como alíquota geral (Decreto 9.393/2018) e **3% para MEI, ME e EPP** exportadores, de 01/08/2025 a 31/12/2026 (Decreto 12.565/2025). O executivo pode citar essas alíquotas legais, mas **não projeta valor de crédito em reunião**: a alíquota aplicável ao caso, o período e o eventual alcance retroativo são confirmados pelo jurídico antes de qualquer projeção. **PARÂMETRO A VALIDAR** — sem exceção. Atenção à janela: o REINTEGRA na forma atual é extinto em janeiro de 2027, com a entrada da CBS.')
question(21, 'Dá para recuperar períodos passados?')
answer('A possibilidade de aproveitar competências anteriores depende da situação da empresa e das regras vigentes em cada período. É análise documental, feita no diagnóstico — o executivo apresenta a **tese**, não o valor. Referência de trabalho: revisão dos últimos **5 anos** de exportações sem crédito apurado, via PER/DCOMP (**PARÂMETRO A VALIDAR**).')

# ---------- BLOCO 5 ----------
H1('BLOCO 5 — JORNADA E OPERAÇÃO')
question(22, 'Quais são as etapas da Rota Fiscal?')
answer('Seis, cada uma com responsável, prazo e evidência documental:')
bullet('**1. Diagnóstico** — raio-X tributário e patrimonial: situação na Receita, DIRPF, estrutura societária e objetivos.')
bullet('**2. Desenho da rota** — definição da rota (ou combinação), modalidade de residência/estrutura, parecer e cronograma.')
bullet('**3. Documentação** — certidões, autenticações, apostilamento de Haia e tradução juramentada: o dossiê completo antes de qualquer ato.')
bullet('**4. Planejamento da saída** — quando aplicável: DIRPF regularizada e Declaração de Saída Definitiva, sem bitributação nem pendência.')
bullet('**5. Missão presencial** — agenda organizada no destino: migrações, registros e abertura da estrutura, com acompanhamento local.')
bullet('**6. Ativação e manutenção** — RUC/registros ativos, contabilidade local e calendário de obrigações: estrutura viva.')
question(23, 'Quanto tempo leva?')
answer('Depende da rota, do perfil migratório e da documentação do cliente. O **cronograma sai no parecer da etapa 2** — personalizado e por escrito. O executivo não promete prazo fechado em reunião: apresenta o método e a etapa em que o prazo é assumido formalmente.')
question(24, 'O que o cliente precisa apresentar?')
answer('No diagnóstico: situação fiscal na Receita, **DIRPF recente**, estrutura societária e uma conversa franca sobre objetivos (receita, patrimônio, família, plano de vida). A lista documental completa vem na etapa 3, organizada pela equipe — o cliente não "corre atrás" sozinho.')
question(25, 'Como funciona a missão presencial?')
answer('A equipe organiza a **agenda no destino** — migrações, registros, abertura de contas e da estrutura — e acompanha o cliente localmente. É a etapa em que a EAS Paraguay faz diferença visível: o cliente chega com hora marcada, não com um guia turístico.')
question(26, 'E depois de ativado, acaba?')
answer('Não — e este é um ponto comercial importante: a rota inclui **manutenção**. Contabilidade local, calendário de obrigações e conformidade contínua nos dois países. Estrutura abandonada vira passivo; estrutura mantida vira ativo.')

# ---------- BLOCO 6 ----------
H1('BLOCO 6 — OBJEÇÕES FREQUENTES')
question(27, '"Isso é coisa de quem quer fugir do fisco."')
answer('Internacionalizar **não é fugir: é escolher, dentro da lei**, a jurisdição e a estrutura mais eficientes para cada parte do negócio — o mesmo que qualquer multinacional faz. A prova está no método: tudo declarado, nos dois países, com advocacia na condução. Quem quer se esconder não contrata quem formaliza.')
question(28, '"Não vou acabar pagando imposto duas vezes?"')
answer('É exatamente o que o **planejamento da saída** evita: DIRPF regularizada, Declaração de Saída Definitiva quando aplicável e desenho da estrutura à luz das regras brasileiras. Bitributação é sintoma de internacionalização **sem método** — o problema que a Rota Fiscal resolve, não o que ela cria.')
question(29, '"Paraguai é seguro juridicamente?"')
answer('O sistema tributário paraguaio é **simples e estável** — poucas regras, alíquotas baixas, dentro do Mercosul. E o cliente não vai sozinho: a Lucro Ativo tem **estrutura própria em Assunção**, com equipe local e rede credenciada. A segurança vem da formalização e da presença — não de promessa.')
question(30, '"Posso fazer isso sozinho ou com um despachante."')
answer('Pode — e é assim que nascem os problemas clássicos: residência "fantasma", CBE omitido, DIRPF inconsistente, estrutura aberta e abandonada. A Rota Fiscal é **advocacia tributária especializada** conduzindo cada etapa com evidência documental. O barato do despachante custa caro na primeira fiscalização.')
question(31, '"Quanto custa?"')
answer('Honorários e custos são apresentados **após o diagnóstico**, porque dependem da rota e do desenho — número antes disso seria chute, e a casa não chuta. A conversa certa é outra: comparar o investimento com o **custo de permanecer como está**, entregando ao fisco a margem que poderia virar reinvestimento. **PARÂMETRO A VALIDAR** antes de qualquer proposta.')
question(32, '"Com a Reforma Tributária, não é melhor esperar?"')
answer('Ao contrário: a transição (2026–2033) impõe **dois sistemas em paralelo** — CBS/IBS convivendo com o modelo atual — elevando o custo de conformidade justamente para quem produz. Esperar é aceitar anos de complexidade dobrada com carga concentrada numa única jurisdição. Diversificar é a resposta racional ao cenário, não uma aposta contra ele.')

# ---------- BASE NORMATIVA ----------
H1('BASE NORMATIVA DE REFERÊNCIA')
para('Âncoras que o executivo pode citar — sem improvisar detalhe de rito. A leitura fina é do jurídico.')
bullet('**Lei 7.547/2025 e Decreto 5.714/2026 (Paraguai)** — novo regime de maquila (bens e serviços): tributo único de 1% sobre o valor agregado local ou sobre a fatura de exportação; substitui a Lei 1.064/97. Tramitação digital no CNIME (Resolução SE-CNIME 14/2026).')
bullet('**Decreto 9.393/2018 e Decreto 12.565/2025** — alíquotas do REINTEGRA: 0,1% geral; 3% para MEI, ME e EPP entre 01/08/2025 e 31/12/2026 (Programa Acredita Exportação, LC 216/2025).')
bullet('**Lei 13.043/2014, arts. 21 a 29** — institui o REINTEGRA: devolução de resíduo tributário ao exportador de bens manufaturados; percentual fixado por decreto.')
bullet('**Lei 14.754/2023** — tributação brasileira de investimentos e entidades controladas no exterior (base do desenho da LLC declarada): regra geral de 15% sobre o lucro da controlada em 31/12 de cada ano.')
bullet('**Solução de Consulta COSIT 56/2026** — enquadra a LLC transparente com sócios não residentes nos EUA como regime fiscal privilegiado (IN RFB 1.037/2010, art. 2º, VII).')
bullet('**Lei 14.596/2023** — novas regras brasileiras de preços de transferência (padrão OCDE), aplicáveis às operações entre a empresa brasileira e a estrutura no exterior.')
bullet('**Declaração de Saída Definitiva do País (DSDP)** — procedimento da Receita Federal para o encerramento da residência fiscal brasileira.')
bullet('**Declaração de Capitais Brasileiros no Exterior (CBE/Bacen)** — obrigação declaratória para residentes com ativos no exterior acima dos limites regulamentares.')
bullet('**EC 132/2023** — Reforma Tributária do consumo: transição 2026–2033 com convivência entre CBS/IBS e o sistema atual.')
doc.add_paragraph()
callout('▸ LEMBRETE FINAL',
        'Este FAQ é **material interno de capacitação**. Nenhum trecho deve ser encaminhado a cliente como está. As versões com foco no cliente saem **após validação técnica/jurídica** dos parâmetros marcados neste guia.')

doc.save('FAQ_Capacitacao_Comercial_RotaFiscal_LucroAtivo_v1.docx')
print('OK')
